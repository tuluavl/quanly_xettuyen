import copy
import io
import os
import platform
import re
import shutil
import sys
import subprocess
import tempfile
import openpyxl
import urllib.parse
try:
    import pythoncom
except ImportError:
    pythoncom = None

import unicodedata
import threading
import zipfile
import uuid
import secrets
import pyotp
import random
import pandas as pd

from django.core.mail import EmailMultiAlternatives, send_mail
from django.core.signing import Signer, BadSignature
from django.core.cache import cache
from concurrent.futures import ThreadPoolExecutor, as_completed
from docx import Document
from lxml import etree
from datetime import datetime, date
try:
    import pythoncom
    from docx2pdf import convert
except (ImportError, ModuleNotFoundError):
    pythoncom = None
    convert = None

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse, JsonResponse, Http404
from django.contrib import messages
from django.db import connection, transaction
from django.db.models import Q, Count, Subquery, OuterRef
from django.core.paginator import Paginator
from django.conf import settings
from django.views.decorators.http import require_POST
from django.contrib.auth.hashers import check_password, make_password
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from .models import CustomUser, Role, RolePermission, Permission, ThiSinhData, KetQuaLocAo, ToHopMon, TruongTHPT, DiemThiVsat, MauImportGiayBao, CapNhatThongTinTrungTuyen, CauHinhGiayBao, AuditLog, UserProfile
from .decorators import custom_login_required, check_permission
from django.contrib.auth.decorators import login_required


from django.contrib.auth import get_user_model
User = get_user_model()  # Tự động lấy CustomUser


        
def send_otp_email_async(subject, text_content, recipient_email, html_content=None):
    """
    Hàm phụ trợ gửi Email bất đồng bộ (Threading).
    Hỗ trợ cả nội dung Plain Text và HTML.
    """
    try:
        msg = EmailMultiAlternatives(
            subject=subject,
            body=text_content,
            from_email=settings.DEFAULT_FROM_EMAIL,
            to=[recipient_email]
        )
        if html_content:
            msg.attach_alternative(html_content, "text/html")
        msg.send(fail_silently=True)
    except Exception:
        pass


def login_view(request):
    # 1. Nếu đã có Session thì chuyển thẳng tới trang chủ (tránh lặp form login)
    if request.session.get('user_id'):
        return redirect('index')

    if request.method == 'POST':
        username_input = request.POST.get('username', '').strip()
        password_input = request.POST.get('password', '').strip()

        # Kiểm tra dữ liệu đầu vào trống
        if not username_input or not password_input:
            messages.error(request, 'Vui lòng nhập đầy đủ tên đăng nhập và mật khẩu!')
            return render(request, 'xettuyen/login.html')

        # Truy vấn user và lấy kèm theo cả thông tin role và profile
        user = CustomUser.objects.filter(username__iexact=username_input).select_related('role', 'profile').first()

        # Kiểm tra mật khẩu (ĐÃ SỬA: Loại bỏ so sánh chuỗi thuần, chỉ dùng check_password)
        is_valid_password = False
        if user and check_password(password_input, user.password):
            is_valid_password = True

        if is_valid_password:
            # Kiểm tra trạng thái tài khoản bị khóa
            is_locked = getattr(user, 'is_locked', False) or (hasattr(user, 'is_active') and not user.is_active)
            if is_locked:
                messages.error(request, 'Tài khoản của bạn đã bị khóa. Vui lòng liên hệ quản trị viên!')
                return render(request, 'xettuyen/login.html')

            # --- XỬ LÝ XÁC THỰC 2 YẾU TỐ (2FA / OTP) ---
            is_2fa_enabled = hasattr(user, 'profile') and user.profile and getattr(user.profile, 'is_2fa_enabled', False)

            if is_2fa_enabled:
                # Kiểm tra xem tài khoản đã cập nhật email chưa
                if not getattr(user, 'email', None):
                    messages.error(request, 'Tài khoản của bạn chưa đăng ký email để nhận mã OTP. Vui lòng liên hệ quản trị viên!')
                    return render(request, 'xettuyen/login.html')

                # 1. Tạo mã OTP ngẫu nhiên 6 chữ số
                otp_code = str(random.randint(100000, 999999))

                # 2. Xóa session cũ & lưu trạng thái xác thực tạm thời
                request.session.flush()
                request.session['pending_user_id'] = user.id
                request.session['otp_code'] = otp_code

                # 3. Chuẩn bị nội dung Email
                text_content = f'Đây là email tự động từ hệ thống gửi mã OTP xác thực đăng nhập của bạn là: {otp_code}\n\nVui lòng không chia sẻ mã này cho bất kỳ ai và không trả lời email này.'

                html_content = f'''
                <div style="font-family: Arial, sans-serif; line-height: 1.6; color: #333;">
                    <p>Đây là email tự động từ hệ thống gửi mã OTP. Mã xác thực đăng nhập của bạn là:</p>
                    <div style="margin: 15px 0;">
                        <span style="font-size: 28px; font-weight: bold; color: #198754; background-color: #e8f5e9; padding: 8px 16px; border-radius: 6px; letter-spacing: 4px; display: inline-block;">
                            {otp_code}
                        </span>
                    </div>
                    <p style="color: #6c757d; font-size: 13px;">Vui lòng không chia sẻ mã này cho bất kỳ ai và không trả lời email này.</p>
                </div>
                '''

                # 4. Kích hoạt luồng gửi mail ngầm (Thread) không làm hoãn tiến trình chuyển trang
                threading.Thread(
                    target=send_otp_email_async,
                    args=(
                        '[Hệ Thống Xét Tuyển UEH Mekong] Mã xác thực đăng nhập (OTP)',
                        text_content,
                        user.email,
                        html_content
                    )
                ).start()

                messages.info(request, f'Mã OTP đang được gửi đến email {user.email}. Vui lòng kiểm tra hòm thư và nhập mã!')
                return redirect('verify_otp')
            # --- KẾT THÚC XỬ LÝ 2FA ---

            # Đăng nhập trực tiếp nếu 2FA đang TẮT
            request.session.flush()

            # 1. Lưu thông tin tài khoản
            request.session['user_id'] = user.id
            request.session['username'] = user.username
            request.session['full_name'] = getattr(user, 'full_name', '') or user.username

            # 2. Lưu Vai trò & Quyền hạn
            if user.role:
                request.session['role_id'] = user.role.id
                request.session['role_name'] = user.role.name
                request.session['role_code'] = user.role.code
                request.session['permissions'] = list(user.role.permissions.values_list('code', flat=True))
            else:
                request.session['role_id'] = None
                request.session['role_name'] = ''
                request.session['role_code'] = ''
                request.session['permissions'] = []

            messages.success(request, f'Đăng nhập thành công! Chào mừng {request.session["full_name"]}.')
            return redirect('index')

        messages.error(request, 'Tên đăng nhập hoặc mật khẩu không đúng!')

    return render(request, 'xettuyen/login.html')

# XỬ LÝ OTP
def verify_otp_view(request):
    # Lấy thông tin user chờ xác thực từ session tạm
    pending_user_id = request.session.get('pending_user_id')
    saved_otp = request.session.get('otp_code')

    # Nếu truy cập thẳng mà chưa qua bước đăng nhập mật khẩu -> Đẩy về trang login
    if not pending_user_id or not saved_otp:
        messages.error(request, 'Phiên xác thực không hợp lệ hoặc đã hết hạn!')
        return redirect('login')

    if request.method == 'POST':
        otp_input = request.POST.get('otp', '').strip()

        # Kiểm tra mã OTP người dùng nhập vào
        if otp_input == saved_otp:
            user = CustomUser.objects.filter(id=pending_user_id).select_related('role').first()

            if not user:
                messages.error(request, 'Tài khoản không tồn tại!')
                return redirect('login')

            # Xóa các biến session tạm dùng cho 2FA
            del request.session['pending_user_id']
            del request.session['otp_code']

            # --- THIẾT LẬP SESSION ĐĂNG NHẬP CHÍNH THỨC ---
            request.session['user_id'] = user.id
            request.session['username'] = user.username
            request.session['full_name'] = user.full_name or user.username

            if user.role:
                request.session['role_id'] = user.role.id
                request.session['role_name'] = user.role.name
                request.session['role_code'] = user.role.code
                request.session['permissions'] = list(user.role.permissions.values_list('code', flat=True))
            else:
                request.session['role_id'] = None
                request.session['role_name'] = ''
                request.session['role_code'] = ''
                request.session['permissions'] = []

            messages.success(request, f'Xác thực thành công! Chào mừng {user.full_name or user.username}.')
            return redirect('index')

        # Trường hợp nhập sai OTP
        messages.error(request, 'Mã OTP không chính xác. Vui lòng thử lại!')

    return render(request, 'xettuyen/verify_otp.html')
    
def logout_view(request):
    # 1. Xóa toàn bộ dữ liệu Session hiện tại
    request.session.flush()
    
    # 2. Tạo thông báo (messages sẽ lưu vào session mới chỉ dành riêng cho việc hiển thị ở login)
    messages.info(request, 'Đã đăng xuất tài khoản thành công!')
    
    # 3. Chuyển hướng trực tiếp về trang đăng nhập
    return redirect('login')

#PHÂN HỆ INDEX VÀ DANH SÁCH THÍ SINH
@custom_login_required
def index(request):
    role_code = request.session.get('role_code', '')
    username = request.session.get('username', '')
    permissions = request.session.get('permissions', [])

    is_admin = (role_code == 'admin' or username == 'admin')

    PERMISSION_REDIRECT_MAP = {
        'diem_chuan': 'diem_chuan',
        'danh_sach_diem_vsat': 'danh_sach_diem_vsat',
        'danh_sach_to_hop_mon': 'to_hop_mon',
        'import_diem_vsat': 'import_diem_vsat',
        'danh_sach_truong_thpt': 'danh_sach_truong_thpt',
    }

    if is_admin or 'danh_sach_thi_sinh' in permissions:
        qs = ThiSinhData.objects.all()
        tong_thi_sinh = qs.count()

        # 1. Thống kê Nguyện vọng (Loại bỏ null, rỗng, nan, 0 và '0')
        invalid_nv_condition = lambda field: (
            Q(**{f"{field}__isnull": True}) | 
            Q(**{field: ''}) | 
            Q(**{f"{field}__iexact": 'nan'}) | 
            Q(**{field: '0'}) | 
            Q(**{field: 0})
        )

        count_nv1 = qs.exclude(invalid_nv_condition('nv1')).count()
        count_nv2 = qs.exclude(invalid_nv_condition('nv2')).count()
        count_nv3 = qs.exclude(invalid_nv_condition('nv3')).count()

        # 2. Thống kê Thí sinh đăng ký xét Luật
        qs_khong_dk_luat = qs.filter(
            Q(thi_sinh_xet_luat__isnull=True) | 
            Q(thi_sinh_xet_luat='') | 
            Q(thi_sinh_xet_luat__icontains='Không ĐK') | 
            Q(thi_sinh_xet_luat__icontains='Khong DK')
        )
        xet_luat_khong_dk = qs_khong_dk_luat.count()
        xet_luat_co_dk = tong_thi_sinh - xet_luat_khong_dk

        # 3. Thống kê Điều kiện xét Luật (chỉ tính trong nhóm CÓ ĐĂNG KÝ XÉT LUẬT)
        qs_co_dk_luat = qs.exclude(
            Q(thi_sinh_xet_luat__isnull=True) | 
            Q(thi_sinh_xet_luat='') | 
            Q(thi_sinh_xet_luat__icontains='Không ĐK') | 
            Q(thi_sinh_xet_luat__icontains='Khong DK')
        )
        
        luat_du_dieu_kien = qs_co_dk_luat.filter(
            dieu_kien_nganh_luat__icontains='Đủ điều kiện'
        ).exclude(dieu_kien_nganh_luat__icontains='Không').count()
        
        # Không đủ điều kiện = Tổng số thí sinh có ĐK xét Luật - Số thí sinh đủ điều kiện
        luat_khong_du_dieu_kien = max(0, xet_luat_co_dk - luat_du_dieu_kien)

        # 4. Thống kê Phương thức xét tuyển (PTXT)
        ptxt_dgnl = qs.filter(diem_DGNL__gt=0).count()
        ptxt_vsat = qs.filter(diem_VSAT__gt=0).count()
        ptxt_thpt = qs.filter(
            (Q(diem_DGNL__isnull=True) | Q(diem_DGNL=0)) & 
            (Q(diem_VSAT__isnull=True) | Q(diem_VSAT=0))
        ).count()

        headers = [field.name for field in ThiSinhData._meta.fields]
        danh_sach = list(qs.values())

        context = {
            'headers': headers,
            'danh_sach': danh_sach,
            'tong_so': tong_thi_sinh,
            'count_nv1': count_nv1,
            'count_nv2': count_nv2,
            'count_nv3': count_nv3,
            'xet_luat_co_dk': xet_luat_co_dk,
            'xet_luat_khong_dk': xet_luat_khong_dk,
            'luat_du_dieu_kien': luat_du_dieu_kien,
            'luat_khong_du_dieu_kien': luat_khong_du_dieu_kien,
            'ptxt_dgnl': ptxt_dgnl,
            'ptxt_vsat': ptxt_vsat,
            'ptxt_thpt': ptxt_thpt,
        }
        return render(request, 'xettuyen/index.html', context)

    for perm_code in permissions:
        if perm_code in PERMISSION_REDIRECT_MAP:
            return redirect(PERMISSION_REDIRECT_MAP[perm_code])

    return render(request, 'xettuyen/no_permission.html')
    
    
#LẤY NGÀY THÁNG NĂM KÈM VỚI CÁC FILE XUẤT EXCEL, PDF
def create_file_response(prefix, extension="xlsx", content_type=None):
    """
    Tự động gắn ngày hiện tại (DDMMYYYY) và mã hóa chuẩn Unicode tương thích mọi trình duyệt.
    """
    content_types = {
        'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'zip': 'application/zip',
        'pdf': 'application/pdf',
        'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    }
    
    ext = extension.lstrip('.')
    # Định dạng ngày: 11092026
    today_str = datetime.now().strftime('%d%m%Y')
    
    # Tạo tên file hoàn chỉnh
    filename = f"{prefix}_{today_str}.{ext}"
    encoded_filename = urllib.parse.quote(filename)
    
    c_type = content_type or content_types.get(ext, 'application/octet-stream')
    response = HttpResponse(content_type=c_type)
    
    # Chuẩn Header ghép cả tên file thường và tên file mã hóa UTF-8
    response['Content-Disposition'] = f'attachment; filename="{filename}"; filename*=UTF-8\'\'{encoded_filename}'
    return response
    
#PHÂN HỆ THÍ SINH
# Regex kiểm tra tên cột: Chỉ cho phép chữ cái, chữ số và dấu gạch dưới
SAFE_COLUMN_PATTERN = re.compile(r'^[a-zA-Z0-9_]+$')

@custom_login_required
@check_permission('import_thi_sinh')
def import_thi_sinh(request):
    """Đọc file Excel, xóa sạch dữ liệu cũ và import toàn bộ dữ liệu mới (An toàn SQL Injection)"""
    if request.method == 'POST' and request.FILES.get('excel_file'):
        excel_file = request.FILES['excel_file']

        try:
            # 1. Đọc dữ liệu từ file Excel
            df = pd.read_excel(excel_file, dtype=str)

            if not df.empty:
                columns = [str(col).strip() for col in df.columns]

                # --- CHỐNG SQL INJECTION TÊN CỘT (HEADER VALIDATION) ---
                invalid_cols = [col for col in columns if not SAFE_COLUMN_PATTERN.match(col)]
                if invalid_cols:
                    messages.error(
                        request, 
                        f'Import thất bại! Phát hiện tên cột chứa ký tự không hợp lệ: {", ".join(invalid_cols)}'
                    )
                    return redirect(request.META.get('HTTP_REFERER', '/'))

                cols_formatted = ', '.join([f'`{col}`' for col in columns])
                placeholders = ', '.join(['%s'] * len(columns))

                sql_insert = f"""
                    INSERT INTO `thi_sinh_data` ({cols_formatted}) 
                    VALUES ({placeholders})
                """
                
                # Chuẩn hóa dữ liệu đầu vào (loại bỏ NaN và khoảng trắng)
                values = [
                    tuple(None if pd.isna(val) else str(val).strip() for val in row)
                    for row in df.to_numpy()
                ]

                # 2. Thực thi xóa cũ & thêm mới trong một Transaction
                with transaction.atomic():
                    with connection.cursor() as cursor:
                        # Xóa sạch toàn bộ dữ liệu bảng cũ
                        ThiSinhData.objects.all().delete()
                        
                        # Chèn dữ liệu mới qua Parametric Query (%s) an toàn
                        cursor.executemany(sql_insert, values)

                messages.success(
                    request,
                    f'Đã xóa dữ liệu cũ và Import thành công {len(df)} dòng dữ liệu mới!'
                )
            else:
                messages.warning(request, 'File Excel không có dữ liệu.')

        except Exception as e:
            messages.error(request, f'Lỗi khi import dữ liệu: {str(e)}')

    return redirect(request.META.get('HTTP_REFERER', '/'))


@custom_login_required
@check_permission('export_thi_sinh')
def export_thi_sinh(request):
    """Lấy dữ liệu từ bảng thi_sinh_data và xuất ra file Excel"""
    try:
        query = 'SELECT * FROM thi_sinh_data'
        df = pd.read_sql_query(query, connection)

        # Trả về tên file: Danh_sach_thi_sinh_11092026.xlsx
        response = create_file_response('Danh_sach_thi_sinh', extension='xlsx')

        with pd.ExcelWriter(response, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='ThiSinhData')

        return response

    except Exception as e:
        messages.error(request, f'Lỗi khi xuất file Excel: {str(e)}')
        return redirect(request.META.get('HTTP_REFERER', '/'))
    
# ==========================================
# 1. CÁC HÀM HỖ TRỢ BẢO MẬT & AUDIT LOGS
# ==========================================

def generate_2fa_secret():
    """Tạo chuỗi bí mật ngẫu nhiên cho user"""
    return pyotp.random_base32()


def verify_2fa_token(secret, token):
    """Kiểm tra mã 6 chữ số từ app Authenticator"""
    totp = pyotp.TOTP(secret)
    return totp.verify(token)


def log_audit(request, action, target_user=None, details=""):
    """Hàm phụ trợ ghi nhận vết thao tác hệ thống vào CSDL"""
    x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    ip = x_forwarded_for.split(',')[0] if x_forwarded_for else request.META.get('REMOTE_ADDR')
    
    # Lấy tài khoản CustomUser đang đăng nhập từ Session
    actor = None
    user_id = request.session.get('user_id')
    if user_id:
        actor = CustomUser.objects.filter(id=user_id).first()

    AuditLog.objects.create(
        actor=actor,
        action=action,
        target_user=target_user,
        details=details,
        ip_address=ip
    )

def force_logout_user(user):
    """Hàm phụ trợ tăng session_version để buộc user đăng xuất trên các thiết bị khác"""
    profile, _ = UserProfile.objects.get_or_create(user=user)
    profile.session_version += 1
    profile.save()
# ==========================================
# 2. VIEW QUẢN LÝ USER & PHÂN QUYỀN
# ==========================================

@custom_login_required
@check_permission('manage_users')
def manage_users(request):
    GROUP_MAP = dict(Permission.GROUP_CHOICES)

    if request.method == 'POST':
        action = request.POST.get('action')

        try:
            # --- 1. TẠO TÀI KHOẢN MỚI ---
            if action == 'create_user':
                username = request.POST.get('username', '').strip()
                full_name = request.POST.get('full_name', '').strip()
                email = request.POST.get('email', '').strip()  # Bổ sung email
                password = request.POST.get('password', '').strip()
                role_id = request.POST.get('role_id')

                if CustomUser.objects.filter(username=username).exists():
                    messages.error(request, f'Tên đăng nhập "{username}" đã tồn tại!')
                else:
                    user = CustomUser(
                        username=username,
                        full_name=full_name,
                        email=email if email else None,  # Cập nhật email
                        role_id=role_id if role_id else None
                    )
                    if password:
                       user.password = make_password(password)
                    
                    user.save()

                    # Khởi tạo Profile & Ghi nhật ký
                    UserProfile.objects.get_or_create(user=user)
                    log_audit(request, 'CREATE_USER', user, f'Khởi tạo tài khoản mới "{username}" (Họ tên: {full_name}, Email: {email})')
                    messages.success(request, f'Đã tạo tài khoản "{username}" thành công!')

            # --- 2. CHỈNH SỬA THÔNG TIN USER ---
            elif action == 'edit_user':
                user_id = request.POST.get('user_id')
                
                try:
                    user = CustomUser.objects.get(id=user_id)
                except CustomUser.DoesNotExist:
                    messages.error(request, 'Tài khoản không tồn tại trên hệ thống!')
                    return redirect('manage_users')

                full_name = request.POST.get('full_name', '').strip()
                email = request.POST.get('email', '').strip()  # Bổ sung email
                role_id = request.POST.get('role_id')
                password = request.POST.get('password', '').strip()

                old_role_id = user.role_id
                user.full_name = full_name
                user.email = email if email else None  # Cập nhật email
                user.role_id = role_id if role_id else None

                pwd_changed = False
                if password:
                    user.password = make_password(password)
                    pwd_changed = True

                user.save()

                # Nếu có đổi mật khẩu hoặc thay đổi Vai trò -> Đăng xuất tài khoản ở các phiên khác
                if pwd_changed or str(old_role_id or '') != str(role_id or ''):
                    force_logout_user(user)
                    log_audit(request, 'UPDATE_USER', user, f'Cập nhật thông tin/mật khẩu và hủy Session hoạt động cũ của "{user.username}"')
                else:
                    log_audit(request, 'UPDATE_USER', user, f'Cập nhật thông tin cá nhân của "{user.username}"')

                messages.success(request, f'Đã cập nhật thông tin cho tài khoản "{user.username}"!')

            # --- 3. XÓA MỀM / VÔ HIỆU HÓA TÀI KHOẢN ---
            elif action == 'delete_user':
                user_id = request.POST.get('user_id')
                user = CustomUser.objects.get(id=user_id)

                if user.username == 'admin':
                    messages.error(request, 'Không thể vô hiệu hóa hoặc xóa tài khoản Admin hệ thống!')
                else:
                    user.is_active = False
                    user.save()
                    
                    # Đăng xuất cưỡng chế ngay lập tức
                    force_logout_user(user)
                    
                    log_audit(request, 'TOGGLE_STATUS', user, f'Đã vô hiệu hóa (Xóa mềm) tài khoản "{user.username}" và cưỡng chế Đăng xuất.')
                    messages.success(request, f'Đã vô hiệu hóa tài khoản "{user.username}" thành công!')

            # --- 4. KHÓA / MỞ KHÓA TÀI KHOẢN (TOGGLE STATUS) ---
            elif action == 'toggle_user_status':
                user_id = request.POST.get('user_id')
                user = CustomUser.objects.get(id=user_id)

                if user.username == 'admin':
                    messages.error(request, 'Không thể thao tác trên tài khoản Admin hệ thống!')
                else:
                    user.is_active = not user.is_active
                    user.save()

                    if not user.is_active:
                        force_logout_user(user)

                    status_str = "Kích hoạt lại" if user.is_active else "Khóa / Vô hiệu hóa"
                    log_audit(request, 'TOGGLE_STATUS', user, f'{status_str} tài khoản "{user.username}"')
                    messages.success(request, f'Đã {status_str} tài khoản "{user.username}"!')

            # --- 5. CẬP NHẬT VAI TRÒ ---
            elif action == 'update_role':
                user_id = request.POST.get('user_id')
                role_id = request.POST.get('role_id')
                user = CustomUser.objects.get(id=user_id)
                user.role_id = role_id if role_id else None
                user.save()

                force_logout_user(user)
                log_audit(request, 'UPDATE_USER', user, f'Thay đổi vai trò người dùng sang Role ID: {role_id}')
                messages.success(request, f'Đã cập nhật vai trò cho {user.username}')

            # --- 6. CẬP NHẬT DANH SÁCH QUYỀN CHO VAI TRÒ ---
            elif action == 'update_permissions':
                role_id = request.POST.get('role_id')
                perm_ids = request.POST.getlist('permissions')
                role = Role.objects.get(id=role_id)
                role.permissions.set(perm_ids)

                affected_users = CustomUser.objects.filter(role_id=role_id)
                for u in affected_users:
                    force_logout_user(u)

                log_audit(request, 'CHANGE_PERM', None, f'Cập nhật lại danh sách quyền hạn cho Vai trò "{role.name}". Đã hủy session của {affected_users.count()} cán bộ liên quan.')
                messages.success(request, f'Đã cập nhật danh sách quyền cho vai trò {role.name}')

            # --- 7. TẠO QUYỀN MỚI ---
            elif action == 'create_permission':
                perm_name = request.POST.get('perm_name', '').strip()
                perm_code = request.POST.get('perm_code', '').strip()
                perm_group = request.POST.get('perm_group', 'other')
                
                if Permission.objects.filter(code=perm_code).exists():
                    messages.error(request, f'Mã quyền "{perm_code}" đã tồn tại!')
                else:
                    perm = Permission.objects.create(name=perm_name, code=perm_code, group=perm_group)
                    log_audit(request, 'CHANGE_PERM', None, f'Tạo quyền mới trên hệ thống: "{perm_name}" ({perm_code})')
                    messages.success(request, f'Đã tạo thành công quyền mới: {perm_name}')

            # --- 8. CHỈNH SỬA QUYỀN ---
            elif action == 'edit_permission':
                perm_id = request.POST.get('perm_id')
                perm_name = request.POST.get('perm_name', '').strip()
                perm_code = request.POST.get('perm_code', '').strip()
                perm_group = request.POST.get('perm_group')

                perm = Permission.objects.get(id=perm_id)
                if Permission.objects.filter(code=perm_code).exclude(id=perm_id).exists():
                    msg = f'Mã quyền "{perm_code}" đã tồn tại ở quyền khác!'
                    if request.POST.get('is_ajax') == '1' or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                        return JsonResponse({'status': 'error', 'message': msg}, status=400)
                    messages.error(request, msg)
                else:
                    perm.name = perm_name
                    perm.code = perm_code
                    if perm_group:
                        perm.group = perm_group
                    perm.save()

                    log_audit(request, 'CHANGE_PERM', None, f'Chỉnh sửa thông tin quyền ID {perm.id}: "{perm_name}" ({perm_code})')

                    msg = f'Đã cập nhật quyền "{perm.name}" thành công!'
                    if request.POST.get('is_ajax') == '1' or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                        return JsonResponse({
                            'status': 'success',
                            'message': msg,
                            'perm_id': perm.id,
                            'perm_name': perm.name,
                            'perm_code': perm.code,
                            'perm_group_label': GROUP_MAP.get(perm.group, perm.group)
                        })
                    messages.success(request, msg)

            # --- 9. XÓA QUYỀN ---
            elif action == 'delete_permission':
                perm_id = request.POST.get('perm_id')
                perm = Permission.objects.get(id=perm_id)
                perm_name = perm.name
                perm_code = perm.code
                perm.delete()

                log_audit(request, 'CHANGE_PERM', None, f'Đã xóa quyền "{perm_name}" ({perm_code}) khỏi hệ thống.')

                msg = f'Đã xóa quyền "{perm_name}" thành công!'
                if request.POST.get('is_ajax') == '1' or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({'status': 'success', 'message': msg, 'perm_id': perm_id})
                messages.success(request, msg)

            # --- 10. CHUYỂN NHÓM CHO 1 QUYỀN ---
            elif action == 'update_permission_group':
                perm_id = request.POST.get('perm_id')
                new_group = request.POST.get('new_group')
                perm = Permission.objects.get(id=perm_id)
                perm.group = new_group
                perm.save()

                log_audit(request, 'CHANGE_PERM', None, f'Đã chuyển quyền "{perm.name}" sang nhóm "{new_group}"')
                messages.success(request, f'Đã chuyển quyền "{perm.name}" sang nhóm mới!')
            # --- 12. BẬT TẮT XÁC THỰC 2FA (UPDATE) ---
            elif action == 'toggle_2fa':
                user_id = request.POST.get('user_id')
                user = CustomUser.objects.get(id=user_id)
                profile, _ = UserProfile.objects.get_or_create(user=user)
    
                profile.is_2fa_enabled = not profile.is_2fa_enabled
                profile.save()

                status_str = "Bật" if profile.is_2fa_enabled else "Tắt"
                messages.success(request, f'Đã {status_str} xác thực 2 bước cho tài khoản "{user.username}"!')
            # --- 12. CHUYỂN NHÓM HÀNG LOẠT (BULK UPDATE) ---
            elif action == 'bulk_update_group':
                perm_ids = request.POST.getlist('perm_ids')
                target_group = request.POST.get('target_group')
                
                if perm_ids and target_group:
                    Permission.objects.filter(id__in=perm_ids).update(group=target_group)
                    target_label = GROUP_MAP.get(target_group, target_group)

                    log_audit(request, 'CHANGE_PERM', None, f'Chuyển hàng loạt {len(perm_ids)} quyền sang nhóm "{target_label}"')

                    if request.POST.get('is_ajax') == '1' or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                        return JsonResponse({
                            'status': 'success',
                            'message': f'Đã chuyển thành công {len(perm_ids)} quyền sang nhóm "{target_label}"!',
                            'updated_ids': perm_ids,
                            'target_group_label': target_label
                        })

                    messages.success(request, f'Đã cập nhật nhóm cho {len(perm_ids)} quyền thành công!')

        except Exception as e:
            if request.POST.get('is_ajax') == '1' or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({'status': 'error', 'message': f'Lỗi hệ thống: {str(e)}'}, status=400)
            messages.error(request, f'Lỗi hệ thống: {str(e)}')
            
        return redirect('manage_users')

   # ==========================================
    # QUERIES DỮ LIỆU HIỂN THỊ
    # ==========================================
    users = CustomUser.objects.select_related('role', 'profile').prefetch_related('role__permissions').all()
    roles = Role.objects.prefetch_related('permissions').all()
    audit_logs = AuditLog.objects.select_related('actor', 'target_user').all()[:100]

    # Bổ sung tự động khởi tạo Profile cho các user bị thiếu
    for u in users:
        if not hasattr(u, 'profile') or u.profile is None:
            UserProfile.objects.get_or_create(user=u)

    # Gom nhóm quyền ĐỘNG trực tiếp theo cột 'group' trong Database
    grouped_permissions = {}

    for group_key, group_label in GROUP_MAP.items():
        if group_key == 'other':
            continue
        
        # Trích xuất chính xác theo tên nhóm trong CSDL (group=group_key)
        perms = Permission.objects.filter(group=group_key)
        
        if perms.exists():
            grouped_permissions[group_label] = perms

    # Gom các quyền thuộc nhóm 'other' hoặc các group chưa định nghĩa trong GROUP_MAP
    other_perms = Permission.objects.filter(Q(group='other') | ~Q(group__in=GROUP_MAP.keys()))
    if other_perms.exists():
        grouped_permissions['Chức năng khác'] = other_perms

    context = {
        'users': users,
        'roles': roles,
        'audit_logs': audit_logs,
        'grouped_permissions': grouped_permissions,
        'group_choices': Permission.GROUP_CHOICES,
        'all_permissions': Permission.objects.all().order_by('group', 'id'),
    }
    return render(request, 'xettuyen/manage_users.html', context)
    
#PHÂN HỆ QUẢN LÝ TỔ HỢP MÔN
@custom_login_required
@check_permission('to_hop_mon')
def to_hop_mon(request):
    with connection.cursor() as cursor:
        cursor.execute("SELECT stt, ma_to_hop_mon, ten_to_hop_mon, ma_mon_thi FROM thm ORDER BY ma_to_hop_mon ASC")
        rows = cursor.fetchall()

    danh_sach = []
    for row in rows:
        ma_mon_raw = str(row[3] or '').strip()
        if ',' in ma_mon_raw:
            mon_list = [m.strip() for m in ma_mon_raw.split(',')]
        elif '-' in ma_mon_raw:
            mon_list = [m.strip() for m in ma_mon_raw.split('-')]
        else:
            mon_list = ma_mon_raw.split()

        danh_sach.append({
            'stt': row[0],
            'ma_to_hop': row[1] or '',
            'ten_to_hop': row[2] or '',
            'mon_1': mon_list[0] if len(mon_list) > 0 else '',
            'mon_2': mon_list[1] if len(mon_list) > 1 else '',
            'mon_3': mon_list[2] if len(mon_list) > 2 else '',
        })

    return render(request, 'xettuyen/to_hop_mon.html', {
        'danh_sach': danh_sach,
        'tong_so': len(danh_sach)
    })


@custom_login_required
@check_permission('import_to_hop_mon')
def import_to_hop_mon(request):
    if request.method == 'POST' and request.FILES.get('excel_file'):
        excel_file = request.FILES['excel_file']
        try:
            wb = openpyxl.load_workbook(excel_file)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))

            if len(rows) > 1:
                objects_to_create = []

                for i, row in enumerate(rows[1:], 1):
                    if not row or not any(row):
                        continue

                    # Đọc chính xác thứ tự 4 cột trong file Excel:
                    # row[0]: STT | row[1]: Mã tổ hợp | row[2]: Tên tổ hợp | row[3]: Mã môn thi
                    stt_val = row[0] if row[0] is not None else i
                    ma_to_hop = str(row[1] or '').strip()
                    ten_to_hop = str(row[2] or '').strip()
                    ma_mon = str(row[3] or '').strip() if len(row) > 3 else ''

                    if ma_to_hop:
                        objects_to_create.append(
                            ToHopMon(
                                stt=stt_val,
                                ma_to_hop_mon=ma_to_hop,
                                ten_to_hop_mon=ten_to_hop,
                                ma_mon_thi=ma_mon,
                            )
                        )

                if objects_to_create:
                    with transaction.atomic():
                        # Xóa toàn bộ dữ liệu danh mục cũ để nạp lại danh mục mới chính xác
                        ToHopMon.objects.all().delete()

                        # Thêm hàng loạt an toàn, tương thích tuyệt đối mọi CSDL
                        ToHopMon.objects.bulk_create(
                            objects_to_create, batch_size=500
                        )

                messages.success(
                    request, 'Import danh sách tổ hợp môn thành công!'
                )
        except Exception as e:
            messages.error(request, f'Lỗi khi import file Excel: {str(e)}')

    return redirect('to_hop_mon')


@custom_login_required
@check_permission('export_to_hop_mon')
def xuat_excel_to_hop_mon(request):
    # Khởi tạo response tự động gắn ngày: "Danh_sach_to_hop_mon_DDMMYYYY.xlsx"
    response = create_file_response('Danh_sach_to_hop_mon', extension='xlsx')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ToHopMon"

    ws.append(['STT', 'Mã tổ hợp môn', 'Tên tổ hợp môn', 'Mã môn thi'])

    with connection.cursor() as cursor:
        cursor.execute("SELECT ma_to_hop_mon, ten_to_hop_mon, ma_mon_thi FROM thm ORDER BY ma_to_hop_mon ASC")
        rows = cursor.fetchall()
        for idx, row in enumerate(rows, 1):
            ws.append([idx, row[0] or '', row[1] or '', row[2] or ''])

    wb.save(response)
    return response

@custom_login_required
@check_permission('sua_to_hop_mon')
def sua_to_hop_mon(request, ma_to_hop):
    if request.method == 'POST':
        ten_to_hop_mon = request.POST.get('ten_to_hop_mon', '').strip()
        ma_mon_thi = request.POST.get('ma_mon_thi', '').strip()

        with connection.cursor() as cursor:
            cursor.execute("""
                UPDATE thm 
                SET ten_to_hop_mon = %s, ma_mon_thi = %s 
                WHERE ma_to_hop_mon = %s
            """, [ten_to_hop_mon, ma_mon_thi, ma_to_hop])

        messages.success(request, f"Đã cập nhật tổ hợp môn {ma_to_hop} thành công!")
    return redirect('to_hop_mon')


@custom_login_required
@check_permission('xoa_to_hop_mon')
def xoa_to_hop_mon(request, ma_to_hop):
    if request.method == 'POST':
        with connection.cursor() as cursor:
            cursor.execute("DELETE FROM thm WHERE ma_to_hop_mon = %s", [ma_to_hop])

        messages.success(request, f"Đã xóa tổ hợp môn {ma_to_hop}!")
    return redirect('to_hop_mon')


def normalize_cccd(val):
    if not val:
        return ''
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s.zfill(12) if len(s) > 0 and len(s) < 12 else s

#QUẢN LÝ ĐIỂM VSAT
@custom_login_required
@check_permission('danh_sach_diem_vsat')
def danh_sach_diem_vsat(request):
    """Hàm hiển thị danh sách và trả dữ liệu Ajax DataTables Server-side"""
    if request.headers.get('x-requested-with') == 'XMLHttpRequest' or request.GET.get('draw'):
        draw = int(request.GET.get('draw', 1))
        start = int(request.GET.get('start', 0))
        length = int(request.GET.get('length', 10))
        search_val = request.GET.get('search[value]', '').strip()

        base_qs = DiemThiVsat.objects.all()
        records_total = base_qs.count()

        queryset = base_qs
        if search_val:
            queryset = queryset.filter(
                Q(so_cccd__icontains=search_val) | Q(ho_ten__icontains=search_val)
            )

        records_filtered = queryset.count()
        data_slice = queryset[start:start + length] if length != -1 else queryset

        data = []
        for idx, item in enumerate(data_slice, start=start + 1):
            data.append({
                'id': item.id,
                'stt': idx,
                'so_cccd': item.so_cccd,
                'ho_ten': item.ho_ten,
                'di_vs': item.di_vs if item.di_vs is not None else '',
                'ho_vs': item.ho_vs if item.ho_vs is not None else '',
                'li_vs': item.li_vs if item.li_vs is not None else '',
                'n1_vs': item.n1_vs if item.n1_vs is not None else '',
                'si_vs': item.si_vs if item.si_vs is not None else '',
                'su_vs': item.su_vs if item.su_vs is not None else '',
                'to_vs': item.to_vs if item.to_vs is not None else '',
                'va_vs': item.va_vs if item.va_vs is not None else '',
                'max_score': item.max_score if item.max_score is not None else '',
                'thmon_a00_vsat': item.thmon_a00_vsat if item.thmon_a00_vsat is not None else '',
                'quy_doi_a00': item.quy_doi_a00 if item.quy_doi_a00 is not None else '',
                'thmon_a01_vsat': item.thmon_a01_vsat if item.thmon_a01_vsat is not None else '',
                'quy_doi_a01': item.quy_doi_a01 if item.quy_doi_a01 is not None else '',
                'thmon_d01_vsat': item.thmon_d01_vsat if item.thmon_d01_vsat is not None else '',
                'quy_doi_d01': item.quy_doi_d01 if item.quy_doi_d01 is not None else '',
                'thmon_d07_vsat': item.thmon_d07_vsat if item.thmon_d07_vsat is not None else '',
                'quy_doi_d07': item.quy_doi_d07 if item.quy_doi_d07 is not None else '',
                'thmon_d09_vsat': item.thmon_d09_vsat if item.thmon_d09_vsat is not None else '',
                'quy_doi_d09': item.quy_doi_d09 if item.quy_doi_d09 is not None else '',
                'thmon_d14_vsat': item.thmon_d14_vsat if item.thmon_d14_vsat is not None else '',
                'quy_doi_d14': item.quy_doi_d14 if item.quy_doi_d14 is not None else '',
                'diem_thi_vsat_max': item.diem_thi_vsat_max if item.diem_thi_vsat_max is not None else '',
            })

        return JsonResponse({
            'draw': draw,
            'recordsTotal': records_total,
            'recordsFiltered': records_filtered,
            'data': data,
        })

    # --- TỐI ƯU CSDL: Gom toàn bộ 7 phép đếm count() riêng lẻ thành 1 câu lệnh SQL duy nhất ---
    stats_to_hop = DiemThiVsat.objects.aggregate(
        total=Count('id'),
        a00=Count('id', filter=Q(thmon_a00_vsat__isnull=False) & ~Q(thmon_a00_vsat=0)),
        a01=Count('id', filter=Q(thmon_a01_vsat__isnull=False) & ~Q(thmon_a01_vsat=0)),
        d01=Count('id', filter=Q(thmon_d01_vsat__isnull=False) & ~Q(thmon_d01_vsat=0)),
        d07=Count('id', filter=Q(thmon_d07_vsat__isnull=False) & ~Q(thmon_d07_vsat=0)),
        d09=Count('id', filter=Q(thmon_d09_vsat__isnull=False) & ~Q(thmon_d09_vsat=0)),
        d14=Count('id', filter=Q(thmon_d14_vsat__isnull=False) & ~Q(thmon_d14_vsat=0)),
    )

    total_count = stats_to_hop['total']

    return render(request, 'xettuyen/diem_vsat.html', {
        'stats_to_hop': stats_to_hop,
        'total_count': total_count,
    })
@custom_login_required
@check_permission('sua_diem_vsat')
def sua_diem_vsat(request, pk):
    """View xử lý lấy thông tin điểm và lưu chỉnh sửa."""
    item = get_object_or_404(DiemThiVsat, pk=pk)
    
    if request.method == 'GET':
        return JsonResponse({
            'id': item.id,
            'so_cccd': item.so_cccd,
            'ho_ten': item.ho_ten,
            'to_vs': item.to_vs,
            'va_vs': item.va_vs,
            'li_vs': item.li_vs,
            'ho_vs': item.ho_vs,
            'si_vs': item.si_vs,
            'su_vs': item.su_vs,
            'di_vs': item.di_vs,
            'n1_vs': item.n1_vs,
            'diem_thi_vsat_max': item.diem_thi_vsat_max,
        })
        
    elif request.method == 'POST':
        try:
            item.so_cccd = request.POST.get('so_cccd', '').strip()
            item.ho_ten = request.POST.get('ho_ten', '').strip()
            item.to_vs = request.POST.get('to_vs') or None
            item.va_vs = request.POST.get('va_vs') or None
            item.li_vs = request.POST.get('li_vs') or None
            item.ho_vs = request.POST.get('ho_vs') or None
            item.si_vs = request.POST.get('si_vs') or None
            item.su_vs = request.POST.get('su_vs') or None
            item.di_vs = request.POST.get('di_vs') or None
            item.n1_vs = request.POST.get('n1_vs') or None
            item.diem_thi_vsat_max = request.POST.get('diem_thi_vsat_max') or None
            item.save()
            return JsonResponse({'success': True, 'message': 'Cập nhật điểm thành công!'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': f'Lỗi: {str(e)}'}, status=400)


@custom_login_required
@check_permission('xoa_diem_vsat')
def xoa_diem_vsat(request, pk):
    """View xử lý xóa dòng điểm thi V-SAT."""
    if request.method == 'POST':
        item = get_object_or_404(DiemThiVsat, pk=pk)
        item.delete()
        return JsonResponse({'success': True, 'message': 'Xóa dữ liệu thành công!'})
    return JsonResponse({'success': False, 'message': 'Phương thức không hợp lệ!'}, status=400)

#QUẢN LÝ ĐIỂM CHUẨN
@custom_login_required
@check_permission('diem_chuan')
def diem_chuan(request):
    if request.method == 'POST':
        action = request.POST.get('action')

        # 1. Cập nhật điểm chuẩn từ form
        for key, val in request.POST.items():
            if key.startswith('diem_dau_'):
                ma_xt = key.replace('diem_dau_', '').strip()
                KetQuaLocAo.objects.filter(
                    Q(ma_xet_tuyen_moi=ma_xt) | Q(ma_nganh=ma_xt)
                ).update(diem_chuan=val)

        # 2. Xử lý tính toán xét tuyển
        if action == 'calculate':
            diem_chuan_map = {}
            for item in KetQuaLocAo.objects.all():
                ma_code = str(item.ma_xet_tuyen_moi or item.ma_nganh or '').strip()
                if ma_code:
                    try:
                        val_str = str(item.diem_chuan or 0).replace(',', '.').strip()
                        diem_chuan_map[ma_code] = float(val_str)
                    except ValueError:
                        diem_chuan_map[ma_code] = 0.0

            ds_thi_sinh = ThiSinhData.objects.all()
            thi_sinh_cap_nhat = []
            DIEM_LIET_THRESHOLD = 1.0 

            for ts in ds_thi_sinh:
                ts.kq_NV1 = ts.ma_NV1 = ts.nv_NV1 = None
                ts.kq_NV2 = ts.ma_NV2 = ts.nv_NV2 = None
                ts.kq_NV3 = ts.ma_NV3 = ts.nv_NV3 = None
                ts.kq_TT = 'Trượt'
                ts.ma_TT = ts.nv_TT = None

                if hasattr(ts, 'diem_thi_TT'):
                    ts.diem_thi_TT = None

                ky_thi_val = getattr(ts, 'ky_thi', None)

                nv_list = [
                    (1, str(getattr(ts, 'nv1', '') or '').strip(), 'NV1', getattr(ts, 'tong_diem_qd_e_xet_nv1', None)),
                    (2, str(getattr(ts, 'nv2', '') or '').strip(), 'NV2', getattr(ts, 'tong_diem_qd_e_xet_nv2', None)),
                    (3, str(getattr(ts, 'nv3', '') or '').strip(), 'NV3', getattr(ts, 'tong_diem_qd_e_xet_nv3', None)),
                ]

                diem_chung_raw = getattr(ts, 'diem_xet_tuyen', None) or getattr(ts, 'tong_diem', None) or getattr(ts, 'diem', 0)

                da_trung_tuyen = False

                for stt, ma_nganh_nv, nv_tag, diem_nv_raw in nv_list:
                    if not ma_nganh_nv:
                        continue

                    # KIỂM TRA ĐIỀU KIỆN RIÊNG NGÀNH LUẬT (7380107)
                    dk_luat = str(getattr(ts, 'dieu_kien_nganh_luat', '') or '').strip()
                    if ma_nganh_nv == '7380107' and 'không đủ điều kiện' in dk_luat.lower():
                        setattr(ts, f'kq_{nv_tag}', 'Không đủ ĐK xét')
                        continue

                    raw_val = diem_nv_raw if diem_nv_raw is not None else diem_chung_raw
                    try:
                        diem_ts = float(str(raw_val).replace(',', '.').strip())
                    except ValueError:
                        diem_ts = 0.0

                    diem_san = diem_chuan_map.get(ma_nganh_nv, 0.0)

                    is_diem_hop_le = diem_ts > DIEM_LIET_THRESHOLD
                    is_dat_diem_chuan = diem_san > 0 and diem_ts >= diem_san

                    if is_dat_diem_chuan and is_diem_hop_le:
                        setattr(ts, f'kq_{nv_tag}', 'D')
                        setattr(ts, f'ma_{nv_tag}', ma_nganh_nv)
                        setattr(ts, f'nv_{nv_tag}', str(stt))

                        ts.kq_TT = 'D'
                        ts.ma_TT = ma_nganh_nv
                        ts.nv_TT = str(stt)

                        if hasattr(ts, 'diem_thi_TT'):
                            ts.diem_thi_TT = ky_thi_val

                        if hasattr(ts, 'diem_xet_tuyen'):
                            ts.diem_xet_tuyen = diem_ts

                        da_trung_tuyen = True
                        break
                    else:
                        setattr(ts, f'kq_{nv_tag}', 'Trượt')

                if not da_trung_tuyen:
                    ts.kq_TT = 'Trượt'
                    ts.ma_TT = ts.nv_TT = None
                    if hasattr(ts, 'diem_thi_TT'):
                        ts.diem_thi_TT = None

                thi_sinh_cap_nhat.append(ts)

            existing_fields = [f.name for f in ThiSinhData._meta.fields]
            target_fields = [
                'kq_TT', 'ma_TT', 'nv_TT',
                'kq_NV1', 'ma_NV1', 'nv_NV1',
                'kq_NV2', 'ma_NV2', 'nv_NV2',
                'kq_NV3', 'ma_NV3', 'nv_NV3',
                'diem_xet_tuyen', 'diem_thi_TT',
            ]
            valid_update_fields = [f for f in target_fields if f in existing_fields]

            if thi_sinh_cap_nhat and valid_update_fields:
                with transaction.atomic():
                    ThiSinhData.objects.bulk_update(thi_sinh_cap_nhat, valid_update_fields)

            messages.success(request, 'Đã tính toán xét tuyển và cập nhật số liệu thành công!')

        # 3. Đồng bộ dữ liệu sang bảng mẫu in giấy báo
        elif action == 'import_3_tables':
            try:
                def parse_num(val):
                    if val is None or val == '':
                        return None
                    try:
                        return float(str(val).replace(',', '.'))
                    except (ValueError, TypeError):
                        return None

                # BƯỚC 1: Lấy danh sách thí sinh trúng tuyển hiện tại
                ds_trung_tuyen = ThiSinhData.objects.filter(kq_TT__iexact='D')
                
                current_tt_cccds = set()
                for ts in ds_trung_tuyen:
                    cccd_raw = str(getattr(ts, 'cccd', '') or getattr(ts, 'CCCD', '') or '').strip()
                    if cccd_raw:
                        current_tt_cccds.add(normalize_cccd(cccd_raw))

                to_delete_ids = []
                for obj in CapNhatThongTinTrungTuyen.objects.all():
                    c_norm = normalize_cccd(obj.cccd) if obj.cccd else None
                    if not c_norm or c_norm not in current_tt_cccds:
                        to_delete_ids.append(obj.id)

                if to_delete_ids:
                    CapNhatThongTinTrungTuyen.objects.filter(id__in=to_delete_ids).delete()

                existing_cap_nhat = {
                    normalize_cccd(obj.cccd): obj 
                    for obj in CapNhatThongTinTrungTuyen.objects.all() 
                    if obj.cccd
                }
                
                to_create_cntt = []
                to_update_cntt = []
                
                for ts in ds_trung_tuyen:
                    cccd_raw = str(getattr(ts, 'cccd', '') or getattr(ts, 'CCCD', '') or '').strip()
                    if not cccd_raw:
                        continue
                    cccd_norm = normalize_cccd(cccd_raw)
                    
                    ccta_val = getattr(ts, 'ccta', None) if hasattr(ts, 'ccta') else getattr(ts, 'CCTA', None)
                    so_diem_ccqt_val = str(ccta_val).strip() if ccta_val is not None and str(ccta_val).strip() != '' else None
                    
                    # Trích xuất 2 trường bổ sung theo yêu cầu
                    tong_diem_val = parse_num(
                        getattr(ts, 'diem_xet_tuyen', None)
                        or getattr(ts, 'diem_xt', None)
                    )

                    dtc0_pt2_val = parse_num(
                        getattr(ts, 'diem_thi_hoc_ba_diem_cong', None)
                        or getattr(ts, 'diem_thi_Hoc_ba_Diem_cong', None)
                        or getattr(ts, 'diem_thi_hocba_diemcong', None)
                        or getattr(ts, 'diem_cong', None)
                    )

                    if cccd_norm in existing_cap_nhat:
                        obj = existing_cap_nhat[cccd_norm]
                        obj.so_diem_ccqt = so_diem_ccqt_val
                        obj.tong_diem = tong_diem_val
                        obj.dtc0_pt2 = dtc0_pt2_val
                        to_update_cntt.append(obj)
                    else:
                        to_create_cntt.append(CapNhatThongTinTrungTuyen(
                            cccd=cccd_norm,
                            so_diem_ccqt=so_diem_ccqt_val,
                            tong_diem=tong_diem_val,
                            dtc0_pt2=dtc0_pt2_val
                        ))

                if to_create_cntt:
                    CapNhatThongTinTrungTuyen.objects.bulk_create(to_create_cntt, batch_size=500)
                if to_update_cntt:
                    CapNhatThongTinTrungTuyen.objects.bulk_update(
                        to_update_cntt, 
                        ['so_diem_ccqt', 'tong_diem', 'dtc0_pt2'], 
                        batch_size=500
                    )

                cap_nhat_dict = {
                    normalize_cccd(obj.cccd): obj 
                    for obj in CapNhatThongTinTrungTuyen.objects.all() 
                    if obj.cccd
                }

                # BƯỚC 2: Chuẩn bị dữ liệu tra cứu liên quan
                thm_dict = {}
                thm_name_dict = {}

                with connection.cursor() as cursor:
                    cursor.execute('SELECT ma_to_hop_mon, ma_mon_thi, ten_to_hop_mon FROM thm')
                    for row in cursor.fetchall():
                        ma_th = str(row[0] or '').strip().upper()
                        ma_mon_raw = str(row[1] or '').strip()
                        ten_th = str(row[2] or '').strip()

                        if ',' in ma_mon_raw:
                            mon_list = [m.strip().upper() for m in ma_mon_raw.split(',')]
                        elif '-' in ma_mon_raw:
                            mon_list = [m.strip().upper() for m in ma_mon_raw.split('-')]
                        else:
                            mon_list = [m.strip().upper() for m in ma_mon_raw.split()]

                        thm_dict[ma_th] = mon_list
                        thm_name_dict[ma_th] = ten_th

                truong_dict = {}
                for tr in TruongTHPT.objects.all():
                    if tr.ma_truong_ghep:
                        truong_dict[tr.ma_truong_ghep.strip()] = tr.ten_truong
                    if tr.ma_truong:
                        truong_dict[tr.ma_truong.strip()] = tr.ten_truong

                vsat_dict = {}
                for v in DiemThiVsat.objects.all():
                    c_key = normalize_cccd(v.so_cccd)
                    if c_key:
                        vsat_dict[c_key] = v

                diem_chuan_dict = {}
                nganh_dict = {}
                for kq in KetQuaLocAo.objects.all():
                    code = str(kq.ma_xet_tuyen_moi or kq.ma_nganh or '').strip()
                    if code:
                        diem_chuan_dict[code] = kq.diem_chuan
                        ten_nganh = getattr(kq, 'ten_ma_xet_tuyen', None) or getattr(kq, 'ten_nganh', None)
                        if ten_nganh:
                            nganh_dict[code] = str(ten_nganh).strip()

                vsat_field_map = {
                    'TO': 'to_vs', 'LI': 'li_vs', 'HO': 'ho_vs', 'VA': 'va_vs',
                    'SU': 'su_vs', 'DI': 'di_vs', 'SI': 'si_vs',
                    'N1': 'n1_vs', 'N2': 'n1_vs', 'N3': 'n1_vs', 'N4': 'n1_vs',
                    'N5': 'n1_vs', 'N6': 'n1_vs', 'N7': 'n1_vs', 'NN': 'n1_vs',
                }

                def get_diem_thpt(ts_obj, mon_code):
                    if not mon_code:
                        return None
                    mon = str(mon_code).strip().upper()
                    if mon in ['N1', 'N2', 'N3', 'N4', 'N5', 'N6', 'N7']:
                        val = getattr(ts_obj, 'nn', None) if getattr(ts_obj, 'nn', None) is not None else getattr(ts_obj, 'NN', None)
                    else:
                        val = getattr(ts_obj, mon.lower(), None) if getattr(ts_obj, mon.lower(), None) is not None else getattr(ts_obj, mon, None)
                    return parse_num(val)

                # BƯỚC 3: Duyệt và xử lý từng thí sinh trúng tuyển
                records_to_insert = []

                for ts in ds_trung_tuyen:
                    cccd_raw = str(getattr(ts, 'cccd', '') or getattr(ts, 'CCCD', '') or '').strip()
                    cccd = normalize_cccd(cccd_raw)

                    cntt_obj = cap_nhat_dict.get(cccd)
                    email_val = cntt_obj.email_sv if cntt_obj and getattr(cntt_obj, 'email_sv', None) else ''
                    ma_sv_val = cntt_obj.mssv if cntt_obj and getattr(cntt_obj, 'mssv', None) else ''

                    ho_ten = str(getattr(ts, 'ho_ten', '') or '').strip()
                    ma_dkxt = getattr(cntt_obj, 'ma_dkxt', None) if cntt_obj and getattr(cntt_obj, 'ma_dkxt', None) else ''
                    ngay_sinh = str(getattr(ts, 'ngay_sinh', '') or '').strip()
                    dtut = str(getattr(ts, 'dtut', '') or '').strip()
                    kvut = str(getattr(ts, 'kvut', '') or '').strip()

                    malop10 = str(getattr(ts, 'malop10', '') or '').strip()
                    malop11 = str(getattr(ts, 'malop11', '') or '').strip()
                    malop12 = str(getattr(ts, 'malop12', '') or '').strip()
                    ten_truong = truong_dict.get(malop12, getattr(ts, 'truong_lop_12', '') or '')
                    hoc_ba = f"{malop10}-{malop11}-{malop12}-{ten_truong}".strip('-')

                    ptxt = str(getattr(ts, 'diem_thi_TT', '') or getattr(ts, 'ky_thi', '') or '').strip()
                    phuong_thuc_xet = 'Xét tuyển tích hợp'

                    ma_nganh_tt = str(getattr(ts, 'ma_TT', '') or '').strip()
                    ctdt = nganh_dict.get(ma_nganh_tt, ma_nganh_tt)

                    nv_tt = str(getattr(ts, 'nv_TT', '') or '').strip()

                    if nv_tt == '1':
                        to_hop_tt = getattr(ts, 'to_hop_nv1', None)
                    elif nv_tt == '2':
                        to_hop_tt = getattr(ts, 'to_hop_nv2', None)
                    elif nv_tt == '3':
                        to_hop_tt = getattr(ts, 'to_hop_nv3', None)
                    else:
                        to_hop_tt = getattr(ts, 'to_hop', None)

                    to_hop_tt = str(to_hop_tt or '').strip().upper()

                    pt2_tn_thm = None
                    pt2_tn_mamon1 = pt2_tn_mamon2 = pt2_tn_mamon3 = None
                    pt2_tn_diemmon1 = pt2_tn_diemmon2 = pt2_tn_diemmon3 = None

                    pt2_vsat_thm = None
                    pt2_vsat_mamon1 = pt2_vsat_mamon2 = pt2_vsat_mamon3 = None
                    pt2_vsat_diemmon1 = pt2_vsat_diemmon2 = pt2_vsat_diemmon3 = None

                    if 'THPT' in ptxt.upper():
                        mon_list = thm_dict.get(to_hop_tt, [])
                        ten_th_db = thm_name_dict.get(to_hop_tt)

                        pt2_tn_thm = f"{to_hop_tt} ({ten_th_db})" if ten_th_db else to_hop_tt

                        if len(mon_list) > 0:
                            pt2_tn_mamon1 = mon_list[0]
                            pt2_tn_diemmon1 = get_diem_thpt(ts, mon_list[0])
                        if len(mon_list) > 1:
                            pt2_tn_mamon2 = mon_list[1]
                            pt2_tn_diemmon2 = get_diem_thpt(ts, mon_list[1])
                        if len(mon_list) > 2:
                            pt2_tn_mamon3 = mon_list[2]
                            pt2_tn_diemmon3 = get_diem_thpt(ts, mon_list[2])

                    if 'VSAT' in ptxt.upper() or 'V-SAT' in ptxt.upper():
                        raw_vsat_code = to_hop_tt.replace('V-SAT', '').replace('VSAT', '').replace('VS', '').strip()
                        mon_list = thm_dict.get(raw_vsat_code, thm_dict.get(to_hop_tt, []))

                        code_vsat = raw_vsat_code or to_hop_tt
                        ten_th_vsat = thm_name_dict.get(raw_vsat_code) or thm_name_dict.get(to_hop_tt)

                        pt2_vsat_thm = f"{code_vsat} ({ten_th_vsat})" if ten_th_vsat else code_vsat

                        vsat_obj = vsat_dict.get(cccd)

                        if len(mon_list) > 0:
                            pt2_vsat_mamon1 = mon_list[0]
                            if vsat_obj:
                                f1 = vsat_field_map.get(mon_list[0])
                                pt2_vsat_diemmon1 = parse_num(getattr(vsat_obj, f1, None)) if f1 else None

                        if len(mon_list) > 1:
                            pt2_vsat_mamon2 = mon_list[1]
                            if vsat_obj:
                                f2 = vsat_field_map.get(mon_list[1])
                                pt2_vsat_diemmon2 = parse_num(getattr(vsat_obj, f2, None)) if f2 else None

                        if len(mon_list) > 2:
                            pt2_vsat_mamon3 = mon_list[2]
                            if vsat_obj:
                                f3 = vsat_field_map.get(mon_list[2])
                                pt2_vsat_diemmon3 = parse_num(getattr(vsat_obj, f3, None)) if f3 else None

                    pt2_dgnl = parse_num(getattr(ts, 'diem_DGNL', None))

                    pt2a_diemtbthpt = parse_num(
                        getattr(ts, 'diem_tb_cac_nam_hoc', None) or getattr(ts, 'Diem_tb_cac_nam_hoc', None)
                    )

                    # PT2_DiemQD: từ diem_thi_hoc_ba
                    pt2_diem_qd = parse_num(
                        getattr(ts, 'diem_thi_hoc_ba', None)
                        or getattr(ts, 'diem_thi_Hoc_ba', None)
                        or getattr(ts, 'diem_thi_hocba', None)
                    )

                    # PT2_QD: từ diem_xet_tuyen
                    pt2_qd = parse_num(
                        getattr(ts, 'diem_xet_tuyen', None)
                        or getattr(ts, 'diem_xt', None)
                    )

                    # DTC0: từ diem_thi_hoc_ba_diem_cong
                    dtc0 = parse_num(
                        getattr(ts, 'diem_thi_hoc_ba_diem_cong', None)
                        or getattr(ts, 'diem_thi_Hoc_ba_Diem_cong', None)
                        or getattr(ts, 'diem_thi_hocba_diemcong', None)
                        or getattr(ts, 'diem_cong', None)
                    )

                    barcode = ''

                    dc = parse_num(diem_chuan_dict.get(ma_nganh_tt, None))

                    cong_hsg = parse_num(getattr(ts, 'cong_giai_hsg', None)) or 0
                    cong_ccta = parse_num(getattr(ts, 'cong_ccta', None)) or 0
                    pt2_diem_cong = None

                    if cong_hsg >= 3:
                        mon_dat_giai = str(getattr(ts, 'mon_dat_giai', '') or getattr(ts, 'mon_giai', '') or '').strip()
                        hang_giai = str(getattr(ts, 'hang', '') or getattr(ts, 'hang_giai', '') or '').strip()
                        pt2_diem_cong = f"- Giải thưởng kỳ thi chọn học sinh giỏi cấp tỉnh/TP: Môn {mon_dat_giai}-Hạng {hang_giai}"
                    elif cong_ccta >= 5:
                        ccta = str(getattr(ts, 'ccta', '') or getattr(ts, 'ccta_val', '') or '').strip()
                        pt2_diem_cong = f"- Chứng chỉ tiếng Anh quốc tế: IELTS {ccta}"

                    records_to_insert.append((
                        cccd, ho_ten, email_val, ma_dkxt, ngay_sinh, dtut, kvut, hoc_ba,
                        ptxt, phuong_thuc_xet, ctdt,
                        pt2_tn_thm, pt2_tn_mamon1, pt2_tn_diemmon1,
                        pt2_tn_mamon2, pt2_tn_diemmon2,
                        pt2_tn_mamon3, pt2_tn_diemmon3,
                        pt2_dgnl,
                        pt2_vsat_thm, pt2_vsat_mamon1, pt2_vsat_diemmon1,
                        pt2_vsat_mamon2, pt2_vsat_diemmon2,
                        pt2_vsat_mamon3, pt2_vsat_diemmon3,
                        pt2a_diemtbthpt, pt2_diem_qd, pt2_qd,
                        ma_sv_val, barcode, dtc0, dc, pt2_diem_cong
                    ))

                # BƯỚC 4: Chèn dữ liệu hàng loạt vào mau_import_giay_bao
                with connection.cursor() as cursor:
                    cursor.execute("DELETE FROM mau_import_giay_bao;")
                    if records_to_insert:
                        sql_insert = """
                            INSERT INTO mau_import_giay_bao (
                                `cccd`, `ho_ten`, `email`, `ma_dkxt`, `ngay_sinh`, `dtut`, `kvut`, `hoc_ba`,
                                `ptxt`, `phuong_thuc_xet`, `ctdt`,
                                `pt2_tn_thm`, `pt2_tn_mamon1`, `pt2_tn_diemmon1`,
                                `pt2_tn_mamon2`, `pt2_tn_diemmon2`,
                                `pt2_tn_mamon3`, `pt2_tn_diemmon3`,
                                `pt2_dgnl`,
                                `pt2_vsat_thm`, `pt2_vsat_mamon1`, `pt2_vsat_diemmon1`,
                                `pt2_vsat_mamon2`, `pt2_vsat_diemmon2`,
                                `pt2_vsat_mamon3`, `pt2_vsat_diemmon3`,
                                `pt2a_diemtbthpt`, `pt2_diem_qd`, `pt2_qd`,
                                `ma_sv`, `barcode`, `dtc0`, `dc`, `pt2_diem_cong`
                            ) VALUES (
                                %s, %s, %s, %s, %s, %s, %s, %s,
                                %s, %s, %s,
                                %s, %s, %s,
                                %s, %s,
                                %s, %s,
                                %s,
                                %s, %s, %s,
                                %s, %s,
                                %s, %s,
                                %s, %s, %s,
                                %s, %s, %s, %s, %s
                            )
                        """
                        cursor.executemany(sql_insert, records_to_insert)

                messages.success(request, f'Đã cập nhật bảng thông tin trúng tuyển và đồng bộ thành công {len(records_to_insert)} thí sinh vào bảng mẫu in giấy báo!')
            except Exception as e:
                messages.error(request, f'Lỗi đồng bộ dữ liệu: {str(e)}')

        else:
            messages.success(request, 'Đã lưu điểm chuẩn thành công!')

    # 4. Thống kê hiển thị ra màn hình
    ds_loc_ao = KetQuaLocAo.objects.all()

    tt_counts = (
        ThiSinhData.objects.filter(kq_TT__iexact='D')
        .values('ma_TT', 'nv_TT')
        .annotate(cnt=Count('pk'))
    )

    count_map = {
        (str(item['ma_TT'] or '').strip(), str(item['nv_TT'] or '').strip()): item['cnt']
        for item in tt_counts
    }

    danh_sach_nganh = []
    loc_ao_updates = []
    tong_chi_tieu = tong_trung_tuyen = 0

    tong_nv1 = 0
    tong_nv2 = 0
    tong_nv3 = 0

    for item in ds_loc_ao:
        ma_code = str(item.ma_xet_tuyen_moi or item.ma_nganh or '').strip()

        cnt_nv1 = count_map.get((ma_code, '1'), 0)
        cnt_nv2 = count_map.get((ma_code, '2'), 0)
        cnt_nv3 = count_map.get((ma_code, '3'), 0)
        cnt_trung_tuyen = cnt_nv1 + cnt_nv2 + cnt_nv3

        tong_nv1 += cnt_nv1
        tong_nv2 += cnt_nv2
        tong_nv3 += cnt_nv3

        try:
            chi_tieu = int(item.chi_tieu_chung or 0)
        except (ValueError, TypeError):
            chi_tieu = 0

        ti_le_vuot = round((cnt_trung_tuyen / chi_tieu) * 100, 2) if chi_tieu > 0 else 0.0

        tong_chi_tieu += chi_tieu
        tong_trung_tuyen += cnt_trung_tuyen

        item.sl_tt = cnt_trung_tuyen
        item.ti_le_tt_ct = ti_le_vuot
        item.nv1 = cnt_nv1
        item.nv2 = cnt_nv2
        item.nv3 = cnt_nv3
        loc_ao_updates.append(item)

        danh_sach_nganh.append({
            'ma_xet_tuyen_moi': item.ma_xet_tuyen_moi,
            'ten_ma_xet_tuyen': item.ten_ma_xet_tuyen,
            'ma_nganh': item.ma_nganh,
            'chi_tieu_chung': chi_tieu,
            'diem_chuan': item.diem_chuan,
            'sl_tt': cnt_trung_tuyen,
            'ti_le_tt_ct': ti_le_vuot,
            'nv1': cnt_nv1,
            'nv2': cnt_nv2,
            'nv3': cnt_nv3,
        })

    if loc_ao_updates:
        with transaction.atomic():
            KetQuaLocAo.objects.bulk_update(
                loc_ao_updates, ['sl_tt', 'ti_le_tt_ct', 'nv1', 'nv2', 'nv3']
            )

    ty_le_vuot_tong = round((tong_trung_tuyen / tong_chi_tieu) * 100, 2) if tong_chi_tieu > 0 else 0.0

    context = {
        'danh_sach_nganh': danh_sach_nganh,
        'tong_chi_tieu': tong_chi_tieu,
        'tong_trung_tuyen': tong_trung_tuyen,
        'ty_le_vuot': ty_le_vuot_tong,
        'tong_nv1': tong_nv1,
        'tong_nv2': tong_nv2,
        'tong_nv3': tong_nv3,
    }
    return render(request, 'xettuyen/diem_chuan.html', context)



@custom_login_required
@check_permission('export_diem_chuan')
def export_diem_chuan(request):
    # Khởi tạo response tự động gắn ngày: "Diem_Chuan_Loc_Ao_DDMMYYYY.xlsx"
    response = create_file_response('Diem_Chuan_Loc_Ao', extension='xlsx')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Diem Chuan"

    headers = [
        "STT", "Mã xét tuyển", "Tên ngành", "Chỉ tiêu", 
        "Điểm chuẩn", "Trúng tuyển", "Tỷ lệ (%)", "NV1", "NV2", "NV3"
    ]
    ws.append(headers)

    danh_sach = KetQuaLocAo.objects.all()

    for idx, item in enumerate(danh_sach, start=1):
        ws.append([
            idx,
            item.ma_xet_tuyen_moi or item.ma_nganh or '',
            item.ten_ma_xet_tuyen or '',
            item.chi_tieu_chung or 0,
            item.diem_chuan or 0,
            item.sl_tt or 0,
            item.ti_le_tt_ct or 0.0,
            item.nv1 or 0,
            item.nv2 or 0,
            item.nv3 or 0
        ])

    wb.save(response)
    return response


# View DANH SÁCH TRÚNG TUYỂN
@custom_login_required
@check_permission('danh_sach_trung_tuyen')
def danh_sach_trung_tuyen(request):
    ds_trung_tuyen = MauImportGiayBao.objects.all()

    stats = {
        'total': ds_trung_tuyen.count(),
        'thpt': ds_trung_tuyen.filter(ptxt__icontains='THPT').count(),
        'vsat': ds_trung_tuyen.filter(Q(ptxt__icontains='VSAT') | Q(ptxt__icontains='V-SAT')).count(),
        'dgnl': ds_trung_tuyen.filter(ptxt__icontains='DGNL').count(),
    }

    context = {
        'danh_sach': ds_trung_tuyen,
        'stats': stats,
    }
    return render(request, 'xettuyen/danh_sach_trung_tuyen.html', context)


# 1. TẠO MASV & BARCODE KHÔNG TRÙNG CHO SINH VIÊN MỚI
@custom_login_required
@check_permission('tao_ma_sv')
def tao_ma_sv(request):
    if request.method == 'POST':
        prefix = request.POST.get('prefix_masv', '').strip()
        if not prefix or len(prefix) != 7 or not prefix.isdigit():
            messages.error(request, "Tiền tố Mã SV phải gồm đúng 7 chữ số!")
            return redirect('danh_sach_trung_tuyen')

        # Bao bọc toàn bộ quá trình tính toán và cập nhật trong Transaction
        with transaction.atomic():
            # 1. Khóa dòng các thí sinh chưa có Mã SV bằng select_for_update()
            # Giúp ngăn chặn các thao tác đồng thời gây trùng lặp mã (Race Condition)
            ts_moi = list(
                MauImportGiayBao.objects.select_for_update().filter(
                    Q(ma_sv__isnull=True) | Q(ma_sv='')
                ).order_by('id')
            )

            total_moi = len(ts_moi)
            if total_moi == 0:
                messages.warning(request, "Tất cả thí sinh trong hệ thống đều đã có Mã SV!")
                return redirect('danh_sach_trung_tuyen')

            # 2. Tìm STT lớn nhất hiện tại theo tiền tố
            max_stt = 0
            existing_masv = MauImportGiayBao.objects.filter(
                ma_sv__startswith=prefix
            ).values_list('ma_sv', flat=True)

            for masv in existing_masv:
                if masv and len(masv) == 11:
                    suffix = masv[7:]
                    if suffix.isdigit():
                        max_stt = max(max_stt, int(suffix))

            # 3. Gán giá trị vào RAM
            stt = max_stt + 1
            for ts in ts_moi:
                ma_sv_moi = f"{prefix}{stt:04d}"
                ts.IDSV = ma_sv_moi
                ts.ma_sv = ma_sv_moi
                ts.barcode = ma_sv_moi
                stt += 1

            # 4. Tối ưu ghi CSDL: Cập nhật hàng loạt (Bulk Update) thay vì ts.save() từng dòng
            MauImportGiayBao.objects.bulk_update(
                ts_moi, 
                ['IDSV', 'ma_sv', 'barcode'], 
                batch_size=1000
            )

        messages.success(
            request, 
            f"Đã tạo tiếp MaSV từ số {max_stt + 1:04d} đến {stt - 1:04d} cho {total_moi} thí sinh mới!"
        )

    return redirect('danh_sach_trung_tuyen')


# 2. CẬP NHẬT SỐ CV & PAGE TỰ ĐỘNG NỐI TIẾP SỐ CŨ
@custom_login_required
@check_permission('cap_nhat_so_cv')
def cap_nhat_so_cv(request):
    if request.method == 'POST':
        raw_cv = request.POST.get('so_cv', '').strip()
        cap_nhat_lai_tat_ca = request.POST.get('mode_all') == '1'

        if cap_nhat_lai_tat_ca:
            danh_sach = MauImportGiayBao.objects.all().order_by('id')
        else:
            danh_sach = MauImportGiayBao.objects.filter(
                Q(so_cv__isnull=True) | Q(so_cv='')
            ).order_by('id')

        total = danh_sach.count()
        if total == 0:
            messages.warning(request, "Không có thí sinh nào cần gán Số CV!")
            return redirect('danh_sach_trung_tuyen')

        if raw_cv.isdigit():
            start_num = int(raw_cv)
        else:
            max_cv = 0
            all_cvs = MauImportGiayBao.objects.exclude(
                Q(so_cv__isnull=True) | Q(so_cv='')
            ).values_list('so_cv', flat=True)

            for cv in all_cvs:
                if str(cv).isdigit():
                    max_cv = max(max_cv, int(cv))

            start_num = (max_cv + 1) if max_cv > 0 else 9900

        with transaction.atomic():
            for index, ts in enumerate(danh_sach):
                current_value = start_num + index
                ts.so_cv = str(current_value)
                ts.page = current_value
                ts.save(update_fields=['so_cv', 'page'])

        messages.success(
            request, 
            f"Đã gán Số CV & Page tiếp theo từ {start_num} đến {start_num + total - 1} cho {total} thí sinh!"
        )
    return redirect('danh_sach_trung_tuyen')


# 3. IMPORT BỔ SUNG THÔNG TIN THÍ SINH
VALID_DB_COLUMNS = {
    'IDSV',
    'so_cv',
    'ho_ten',
    'email',
    'dien_thoai',
    'ma_dkxt',
    'ngay_sinh',
    'dtut',
    'kvut',
    'hoc_ba',
    'phuong_thuc_xet',
    'ptxt',
    'ctdt',
    'pt2_diem_cong',
    'pt2_tn_thm',
    'pt2_tn_mamon1',
    'pt2_tn_diemmon1',
    'pt2_tn_mamon2',
    'pt2_tn_diemmon2',
    'pt2_tn_mamon3',
    'pt2_tn_diemmon3',
    'pt2_dgnl',
    'pt2_vsat_thm',
    'pt2_vsat_mamon1',
    'pt2_vsat_diemmon1',
    'pt2_vsat_mamon2',
    'pt2_vsat_diemmon2',
    'pt2_vsat_mamon3',
    'pt2_vsat_diemmon3',
    'pt2a_diemtbthpt',
    'pt2_diem_qd',
    'pt2_qd',
    'cccd',
    'ma_sv',
    'barcode',
    'dtc0',
    'dc',
    'page',
}

EXCEL_TO_MAU_GIAY_BAO = {
    'số cv': 'so_cv',
    'so_cv': 'so_cv',
    'họ tên': 'ho_ten',
    'ho_ten': 'ho_ten',
    'email': 'email',
    'điện thoại': 'dien_thoai',
    'số điện thoại': 'dien_thoai',
    'dien_thoai': 'dien_thoai',
    'sdt': 'dien_thoai',
    'sđt': 'dien_thoai',
    'mã đkxt': 'ma_dkxt',
    'ma_dkxt': 'ma_dkxt',
    'ngày sinh': 'ngay_sinh',
    'ngay_sinh': 'ngay_sinh',
    'đtut': 'dtut',
    'kvut': 'kvut',
    'đối tượng xét': 'phuong_thuc_xet',
    'phương thức xét': 'phuong_thuc_xet',
    'ptxt': 'ptxt',
    'ctđt': 'ctdt',
    'ctdt': 'ctdt',
    'Điểm cộng': 'pt2_diem_cong',
    'diem cong': 'pt2_diem_cong',
    'pt2_diem_cong': 'pt2_diem_cong',
    'học bạ': 'hoc_ba',
    'hoc_ba': 'hoc_ba',
    'thm tn': 'pt2_tn_thm',
    'môn 1 (tn)': 'pt2_tn_mamon1',
    'điểm 1': 'pt2_tn_diemmon1',
    'môn 2 (tn)': 'pt2_tn_mamon2',
    'điểm 2': 'pt2_tn_diemmon2',
    'môn 3 (tn)': 'pt2_tn_mamon3',
    'điểm 3': 'pt2_tn_diemmon3',
    'điểm đgnl': 'pt2_dgnl',
    'thm vsat': 'pt2_vsat_thm',
    'môn 1 (vsat)': 'pt2_vsat_mamon1',
    'điểm 1.1': 'pt2_vsat_diemmon1',
    'môn 2 (vsat)': 'pt2_vsat_mamon2',
    'điểm 2.1': 'pt2_vsat_diemmon2',
    'môn 3 (vsat)': 'pt2_vsat_mamon3',
    'điểm 3.1': 'pt2_vsat_diemmon3',
    'điểm tb thpt': 'pt2a_diemtbthpt',
    'điểm qđ 1': 'pt2_diem_qd',
    'điểm qđ 2': 'pt2_qd',
    'điểm cộng 0': 'dtc0',
    'dtc0': 'dtc0',
    'điểm chuẩn': 'dc',
    'cccd': 'cccd',
    'mã sv': 'ma_sv',
    'ma_sv': 'ma_sv',
    'barcode': 'barcode',
    'page': 'page',
}
# 1. CẬP NHẬT NÚT ĐỒNG BỘ DỮ LIỆU BỔ SUNG
@custom_login_required
@check_permission('ds_trung_tuyen')
def dong_bo_thong_tin_trung_tuyen(request):
    """Đồng bộ dữ liệu bổ sung từ CapNhatThongTinTrungTuyen sang MauImportGiayBao qua CCCD"""
    if request.method == 'POST':
        ds_bo_sung = CapNhatThongTinTrungTuyen.objects.all()
        map_bo_sung = {normalize_cccd(item.cccd): item for item in ds_bo_sung if item.cccd}

        if not map_bo_sung:
            messages.warning(request, "Không có dữ liệu bổ sung nào để đồng bộ.")
            return redirect('danh_sach_trung_tuyen')

        ds_goc = MauImportGiayBao.objects.all()
        updated_list = []

        for item in ds_goc:
            cccd_key = normalize_cccd(item.cccd) if item.cccd else ""
            src = map_bo_sung.get(cccd_key)

            if src:
                mssv_val = getattr(src, 'mssv', None) or getattr(src, 'masv', None) or ''
                masv_val = getattr(src, 'masv', None) or mssv_val
                email_val = getattr(src, 'email', None) or getattr(src, 'email_sv', None) or ''
                ma_dkxt_val = getattr(src, 'ma_dkxt', None) or ''

                # Xử lý định dạng ngày sinh về dd/mm/YYYY
                ngay_sinh_raw = getattr(src, 'ngay_sinh', None)
                ngay_sinh_val = ''

                if isinstance(ngay_sinh_raw, (datetime, date)):
                    ngay_sinh_val = ngay_sinh_raw.strftime('%d/%m/%Y')
                elif ngay_sinh_raw:
                    ngay_sinh_str = str(ngay_sinh_raw).strip()
                    for fmt in ('%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%d/%m/%Y', '%d-%m-%Y'):
                        try:
                            dt = datetime.strptime(ngay_sinh_str.split('.')[0], fmt)
                            ngay_sinh_val = dt.strftime('%d/%m/%Y')
                            break
                        except ValueError:
                            pass
                    if not ngay_sinh_val:
                        ngay_sinh_val = ngay_sinh_str

                # Cập nhật chính xác các trường theo yêu cầu
                if masv_val:
                    item.IDSV = masv_val
                if mssv_val:
                    item.ma_sv = mssv_val
                    item.barcode = mssv_val
                    if hasattr(item, 'masv'):
                        item.masv = mssv_val
                if email_val:
                    item.email = email_val
                if ma_dkxt_val:
                    item.ma_dkxt = ma_dkxt_val
                if ngay_sinh_val:
                    item.ngay_sinh = ngay_sinh_val

                updated_list.append(item)

        if updated_list:
            fields_to_update = ['IDSV', 'ma_sv', 'barcode', 'email', 'ma_dkxt', 'ngay_sinh']
            if hasattr(MauImportGiayBao, 'masv'):
                fields_to_update.append('masv')

            MauImportGiayBao.objects.bulk_update(updated_list, list(set(fields_to_update)), batch_size=500)
            messages.success(request, f"Đã đồng bộ thành công dữ liệu bổ sung cho {len(updated_list)} thí sinh!")
        else:
            messages.info(request, "Không có bản ghi nào được cập nhật.")

    return redirect('danh_sach_trung_tuyen')   
 
# 2. FILE 1: IMPORT HỆ THỐNG NHẬP HỌC
@custom_login_required
@check_permission('ds_trung_tuyen')
def export_excel_import_nhap_hoc(request):
    """Xuất File 1: Import hệ thống nhập học (Lấy dữ liệu bổ sung từ CapNhatThongTinTrungTuyen)"""
    qs = MauImportGiayBao.objects.all()
    ds_cap_nhat = CapNhatThongTinTrungTuyen.objects.all()
    cntt_map = {normalize_cccd(x.cccd): x for x in ds_cap_nhat if x.cccd}

    data = []
    now = datetime.now()
    ngay_cap_nhat_str = now.strftime("%d/%m/%Y")
    ngay_cap_nhat_val = now.strftime("")

    for item in qs:
        cccd_key = normalize_cccd(item.cccd)
        cntt = cntt_map.get(cccd_key)

        # Xử lý Ngày sinh
        ngay_sinh_raw = getattr(cntt, 'ngay_sinh', None) or getattr(item, 'ngay_sinh', '')
        if isinstance(ngay_sinh_raw, (datetime, date)):
            ngay_sinh_str = ngay_sinh_raw.strftime("%d/%m/%Y")
            ngay_sinh_val = ngay_sinh_raw.strftime("")
        elif ngay_sinh_raw:
            ngay_sinh_str = str(ngay_sinh_raw)
            ngay_sinh_val = str(ngay_sinh_raw)
        else:
            ngay_sinh_str = ""
            ngay_sinh_val = ""

        # Xử lý Ngày thi CCQT
        ngay_thi_ccqt_raw = getattr(cntt, 'ngay_thi_ccqt', None)
        if isinstance(ngay_thi_ccqt_raw, (datetime, date)):
            ngay_thi_ccqt_str = ngay_thi_ccqt_raw.strftime("%d/%m/%Y")
            ngay_thi_ccqt_val = ngay_thi_ccqt_raw.strftime("")
        elif ngay_thi_ccqt_raw:
            ngay_thi_ccqt_str = str(ngay_thi_ccqt_raw)
            ngay_thi_ccqt_val = str(ngay_thi_ccqt_raw)
        else:
            ngay_thi_ccqt_str = ""
            ngay_thi_ccqt_val = ""

        ma_ho_so = getattr(item, 'ma_dkxt', '') or ''
        ho_ten = getattr(item, 'ho_ten', '') or ''
        
        sbd = getattr(cntt, 'sbd', getattr(cntt, 'sbd', '')) or ''
        ma_noi_sinh = getattr(cntt, 'ma_noi_sinh', None) or getattr(cntt, 'ten_tinh_noi_sinh', None) or getattr(cntt, 'ma_noi_sinh', '')
        gioi_tinh = getattr(cntt, 'gioi_tinh', None) or getattr(item, 'gioi_tinh', '')
        email = getattr(item, 'email', '') or ''
        dien_thoai = getattr(item, 'dien_thoai', '') or ''
        cmnd = getattr(item, 'cccd', '') or ''
        dan_toc_uid = getattr(item, 'dan_toc_uid', None) or 1
        
        mssv = getattr(cntt, 'mssv', None) or getattr(cntt, 'masv', None) or getattr(item, 'ma_sv', getattr(item, 'IDSV', '')) or ''
        mat_khau_online = getattr(cntt, 'mat_khau_online', None) or getattr(item, 'mat_khau_online', '') or ''
        email_ueh = getattr(cntt, 'email_ueh', None) or getattr(item, 'email_ueh', '') or ''
        mat_khau_email = getattr(cntt, 'mat_khau_email', None) or getattr(item, 'mat_khau_email', '') or ''
        
        tong_diem = getattr(cntt, 'tong_diem', None) if getattr(cntt, 'tong_diem', None) is not None else (getattr(item, 'pt2_diem_qd', getattr(item, 'tong_diem', '')) or 0)
        dtc0_pt1 = getattr(cntt, 'dtc0_pt1', None)
        dtc0_pt2 = getattr(cntt, 'dtc0_pt2', None) if getattr(cntt, 'dtc0_pt2', None) is not None else getattr(item, 'dtc0', None)
        
        loai_ccqt_uid = getattr(cntt, 'loai_ccqt_uid', None)
        so_diem_ccqt = getattr(cntt, 'so_diem_ccqt', None)
        id_chung_chi = getattr(cntt, 'id_chung_chi', None) or getattr(cntt, 'idchungchi', None)
        
        # Trích xuất chỉ lấy phần số của MaCTDT
        ma_ctdt_raw = str(getattr(item, 'ctdt', getattr(item, 'ma_ctdt', '')) or '').strip()
        match_ma_ctdt = re.search(r'^\d+', ma_ctdt_raw)
        ma_ctdt = match_ma_ctdt.group(0) if match_ma_ctdt else ma_ctdt_raw

        pttt = getattr(item, 'ptxt', getattr(item, 'pttt', '')) or ''

        row = {
            'MaHoSo': ma_ho_so,
            'SBD': sbd,
            'Hoten': ho_ten,
            'NgayCapNhat': ngay_cap_nhat_str,
            'NgayCapNhat_New': ngay_cap_nhat_val,
            'NgaySinh': ngay_sinh_str,
            'NgaySinh_new': ngay_sinh_val,
            'MaNoiSinh': ma_noi_sinh,
            'GioiTinh': gioi_tinh,
            'Email': email,
            'DienThoai': dien_thoai,
            'CMND': cmnd,
            'CCCD': cmnd,
            'DanTocUid': dan_toc_uid,
            'MSSV': mssv,
            'MatKhauOnline': mat_khau_online,
            'EmailUEH': email_ueh,
            'MatKhauEmail': mat_khau_email,
            'TongDiem': tong_diem,
            'DTC0PT1': dtc0_pt1,
            'DTC0PT2': dtc0_pt2,
            'LoaiCCQTUid': loai_ccqt_uid,
            'SoDiemCCQT': so_diem_ccqt,
            'NgayThiCCQT': ngay_thi_ccqt_str,
            'NgayThiCCQT_new': ngay_thi_ccqt_val,
            'IDChungChi': id_chung_chi,
            'MaCTDT': ma_ctdt,
            'pttt': pttt,
        }
        data.append(row)

    df = pd.DataFrame(data)
    
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Import nhap hoc', index=False)
    output.seek(0)

    response = create_file_response("KSV_import_nhap_hoc", extension="xlsx")
    response.content = output.getvalue()
    return response


# 3. FILE 2: HỒ SƠ TRÚNG TUYỂN CHI TIẾT
@custom_login_required
@check_permission('ds_trung_tuyen')
def export_excel_hoso_trung_tuyen(request):
    """Xuất File 2: Hồ sơ trúng tuyển chi tiết (Lấy các cột thiếu từ CapNhatThongTinTrungTuyen)"""
    header_row_1 = [
        'MSSV', 'Phương thức trúng tuyển', 'Mã ngành', 'Chương trình đào tạo trúng tuyển',
        'Số báo danh thi tốt nghiệp THPT năm 2026', 'Họ tên', 'Ngày sinh', 'Giới tính',
        'Nơi sinh_theo giấy khai sinh', 'Nơi sinh_theo địa danh sau sáp nhập', 'Dân tộc',
        'Số Căn cước/CCCD/ĐDCD', 'Đối tượng ưu tiên', 'PT1_Kỳ thi', 'PT1_Môn đạt giải',
        'PT1_Thứ hạng giải thưởng', 'PT1_Năm đạt giải', 'PT2a_Tên tổ hợp môn thi tốt nghiệp THPT',
        'PT2a_Thi THPT_Tên_Môn 1', 'PT2a_Thi THPT_Điểm_Môn 1', 'PT2a_Thi THPT_Tên_Môn 2',
        'PT2a_Thi THPT_Điểm_Môn 2', 'PT2a_Thi THPT_Tên_Môn 3', 'PT2a_Thi THPT_Điểm_Môn 3',
        'PT2a_Điểm thi ĐGNL ĐHQG-HCM', 'PT2a_Tên giải HSG cấp tỉnh/TP', 'PT2a_Môn đạt giải HSG cấp tỉnh/TP',
        'PT2a_Thứ hạng giải HSG cấp tỉnh/TP', 'PT2a_Năm đạt giải HSG cấp tỉnh/TP',
        'Chứng chỉ TAQT - Xét tuyển - Tên', 'Chứng chỉ TAQT - Xét tuyển - Điểm',
        'PT2a_Điểm trung bình Học bạ lớp 10', 'PT2a_Điểm trung bình Học bạ lớp 11', 'PT2a_Điểm trung bình Học bạ lớp 12'
    ]

    header_row_2 = [
        'MSSV', 'TenPTXT', 'MaNganh', 'TenNganhTT', 'SBD', 'HoVaTen', 'NgaySinh', 'GioiTinh',
        'TenTinhNoiSinh', 'TenTinhNoiSinhHienTai', 'TenDanToc', 'CMND', 'DoiTuongUuTien',
        'PT1KyThi', 'PT1MonDatGiai', 'PT1ThuHangGiaiThuong', 'PT1NamDatGiai', 'MaTHM',
        'PT2aTenMon1', 'PT2aDiemMon1', 'PT2aTenMon2', 'PT2aDiemMon2', 'PT2aTenMon3', 'PT2aDiemMon3',
        'DiemDGNL', 'TenGiaiHSG', 'TenMonHSG', 'TenHangHSG', 'NamHSG', 'LoaiCCQTXT', 'DiemCCTA',
        'DiemTBLop10', 'DiemTBLop11', 'DiemTBLop12'
    ]

    qs = MauImportGiayBao.objects.all()
    ds_cap_nhat = CapNhatThongTinTrungTuyen.objects.all()
    cntt_map = {normalize_cccd(x.cccd): x for x in ds_cap_nhat if x.cccd}

    rows_data = [header_row_2]

    for item in qs:
        cccd_key = normalize_cccd(item.cccd)
        cntt = cntt_map.get(cccd_key)

        ngay_sinh_raw = getattr(cntt, 'ngay_sinh', None) or getattr(item, 'ngay_sinh', '')
        if isinstance(ngay_sinh_raw, (datetime, date)):
            ngay_sinh_str = ngay_sinh_raw.strftime("%d/%m/%Y")
        else:
            ngay_sinh_str = str(ngay_sinh_raw or '')

        gioi_tinh = getattr(cntt, 'gioi_tinh', None) or getattr(item, 'gioi_tinh', '')
        ten_tinh_noi_sinh = getattr(cntt, 'ten_tinh_noi_sinh', None) or getattr(cntt, 'ma_noi_sinh', None) or getattr(item, 'ten_tinh_noi_sinh', '')
        
        ten_giai_hsg = getattr(cntt, 'ten_giai_hsg', None) or getattr(item, 'ten_giai_hsg', '')
        ten_mon_hsg = getattr(cntt, 'ten_mon_hsg', None) or getattr(item, 'ten_mon_hsg', '')
        ten_hang_hsg = getattr(cntt, 'ten_hang_hsg', None) or getattr(item, 'ten_hang_hsg', '')
        nam_hsg = getattr(cntt, 'nam_hsg', None) or getattr(item, 'nam_hsg', '')

        loai_ccqt_xt = getattr(cntt, 'loai_ccqt_xt', None) or getattr(cntt, 'loai_ccqt', None) or getattr(item, 'loai_ccqt', '')
        diem_ccta = getattr(cntt, 'diem_ccta', None) or getattr(cntt, 'so_diem_ccqt', None) or getattr(item, 'diem_ccta', '')

        diem_tb_lop10 = getattr(cntt, 'diem_tb_lop10', None) if getattr(cntt, 'diem_tb_lop10', None) is not None else getattr(item, 'diem_tb_lop10', '')
        diem_tb_lop11 = getattr(cntt, 'diem_tb_lop11', None) if getattr(cntt, 'diem_tb_lop11', None) is not None else getattr(item, 'diem_tb_lop11', '')
        diem_tb_lop12 = getattr(cntt, 'diem_tb_lop12', None) if getattr(cntt, 'diem_tb_lop12', None) is not None else getattr(item, 'diem_tb_lop12', '')

        mssv = getattr(cntt, 'mssv', None) or getattr(cntt, 'masv', None) or getattr(item, 'ma_sv', getattr(item, 'mssv', getattr(item, 'IDSV', ''))) or ''

        # Tách MaNganh và TenNganhTT từ chuỗi gốc dạng "7340122: Thương mại điện tử"
        ctdt_raw = str(getattr(item, 'ctdt', getattr(item, 'ma_ctdt', '')) or '').strip()
        if ':' in ctdt_raw:
            parts = ctdt_raw.split(':', 1)
            ma_nganh = parts[0].strip()
            ten_nganh_default = parts[1].strip()
        else:
            match_code = re.search(r'^\d+', ctdt_raw)
            ma_nganh = match_code.group(0) if match_code else ctdt_raw
            ten_nganh_default = ctdt_raw

        ten_nganh_tt = getattr(item, 'ten_nganh_tt', None) or ten_nganh_default
        if ':' in str(ten_nganh_tt):
            ten_nganh_tt = str(ten_nganh_tt).split(':', 1)[1].strip()

        rows_data.append([
            mssv,
            getattr(item, 'ten_ptxt', getattr(item, 'ptxt', 'Phương thức xét tuyển tích hợp')),
            ma_nganh,
            ten_nganh_tt,
            getattr(cntt, 'sbd', getattr(cntt, 'sbd', '')),
            getattr(item, 'ho_ten', ''),
            ngay_sinh_str,
            gioi_tinh,
            ten_tinh_noi_sinh,
            ten_tinh_noi_sinh,
            getattr(item, 'ten_dan_toc', 'Kinh'),
            getattr(item, 'cccd', getattr(item, 'cmnd', '')),
            getattr(item, 'dtut', getattr(item, 'doi_tuong_uu_tien', '')),
            getattr(item, 'pt1_ky_thi', ''), getattr(item, 'pt1_mon_dat_giai', ''), getattr(item, 'pt1_thu_hang', ''), getattr(item, 'pt1_nam', ''),
            getattr(item, 'pt2_tn_thm', getattr(item, 'ma_thm', '')),
            getattr(item, 'pt2_tn_mamon1', getattr(item, 'pt2a_ten_mon1', '')), getattr(item, 'pt2_tn_diemmon1', getattr(item, 'pt2a_diem_mon1', '')),
            getattr(item, 'pt2_tn_mamon2', getattr(item, 'pt2a_ten_mon2', '')), getattr(item, 'pt2_tn_diemmon2', getattr(item, 'pt2a_diem_mon2', '')),
            getattr(item, 'pt2_tn_mamon3', getattr(item, 'pt2a_ten_mon3', '')), getattr(item, 'pt2_tn_diemmon3', getattr(item, 'pt2a_diem_mon3', '')),
            getattr(item, 'pt2_dgnl', getattr(item, 'diem_dgnl', '')),
            ten_giai_hsg, ten_mon_hsg, ten_hang_hsg, nam_hsg,
            loai_ccqt_xt, diem_ccta,
            diem_tb_lop10, diem_tb_lop11, diem_tb_lop12
        ])

    df = pd.DataFrame(rows_data, columns=header_row_1)

    # Khởi tạo response tự động gắn tên "HoSoTrungTuyen_ChiTiet_ImportNhapHoc_DDMMYYYY.xlsx"
    response = create_file_response('HoSoTrungTuyen_ChiTiet_ImportNhapHoc', extension='xlsx')

    # Ghi trực tiếp Pandas DataFrame vào HttpResponse stream
    with pd.ExcelWriter(response, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='DS import', index=False)

    return response
# =========================================================
# 2. XỬ LÝ DANH SÁCH & IMPORT TRÚNG TUYỂN
# =========================================================

# ------------------------------------------------------------------
# HÀM HỖ TRỢ XỬ LÝ DỮ LIỆU EXCEL
# ------------------------------------------------------------------
# 2. Hoàn thiện hàm get_val bị đứt đoạn ở cuối file
def get_val(row, col_idx, default=""):
    try:
        val = row[col_idx].value
        return str(val).strip() if val is not None else default
    except (IndexError, AttributeError):
        return default


def parse_float(val):
    """Chuyển đổi dữ liệu sang kiểu float an toàn."""
    if val is None:
        return None
    try:
        val_str = str(val).replace(',', '.').strip()
        return float(val_str)
    except (ValueError, TypeError):
        return None


def parse_date(val):
    """Định dạng ngày tháng từ các kiểu dữ liệu Excel sang object Date."""
    if val is None:
        return None
    if isinstance(val, (datetime, pd.Timestamp)):
        return val.date()

    val_str = str(val).strip().split(" ")[0]
    if not val_str or val_str.lower() in ['none', 'nan', 'null', '']:
        return None

    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(val_str, fmt).date()
        except ValueError:
            pass
    return None


# ------------------------------------------------------------------
# VIEW XỬ LÝ
# ------------------------------------------------------------------
@custom_login_required
@check_permission('ds_cap_nhat_trung_tuyen')
def danh_sach_cap_nhat_trung_tuyen(request):
  if request.method == 'POST' and (
      request.FILES.get('file_excel') or request.FILES.get('excel_file')
  ):
    file_excel = request.FILES.get('file_excel') or request.FILES.get(
        'excel_file'
    )
    try:
      df = pd.read_excel(file_excel, sheet_name=0, dtype=str)
      df.columns = df.columns.str.strip()

      dict_trung_tuyen = {}
      for _, row in df.iterrows():
        raw_cccd = (
            get_val(row, 'cccd')
            or get_val(row, 'CCCD')
            or get_val(row, 'Số CCCD')
        )
        if not raw_cccd or raw_cccd.lower() in ['nan', 'none', '']:
          continue

        cccd = (
            raw_cccd.zfill(12)
            if len(raw_cccd) < 12 and raw_cccd.isdigit()
            else raw_cccd
        )

        kwargs = {
            'cccd': cccd,
            'ma_dkxt': get_val(row, 'ma_dkxt') or get_val(row, 'Mã ĐKXT'),
            'sbd': (
                get_val(row, 'sbd')
                or get_val(row, 'SBD')
                or get_val(row, 'Số báo danh')
            ),
            'ma_noi_sinh': (
                get_val(row, 'ma_noi_sinh')
                or get_val(row, 'MaNoiSinh')
                or get_val(row, 'Mã nơi sinh')
            ),
            'ten_giai_hsg': (
                get_val(row, 'ten_giai_hsg')
                or get_val(row, 'TenGiaiHSG')
                or get_val(row, 'Tên giải HSG')
            ),
            'ten_mon_hsg': (
                get_val(row, 'ten_mon_hsg')
                or get_val(row, 'TenMonHSG')
                or get_val(row, 'Tên môn HSG')
            ),
            'ten_hang_hsg': (
                get_val(row, 'ten_hang_hsg')
                or get_val(row, 'TenHangHSG')
                or get_val(row, 'Tên hạng HSG')
            ),
            'nam_hsg': (
                get_val(row, 'nam_hsg')
                or get_val(row, 'NamHSG')
                or get_val(row, 'Năm HSG')
            ),
            'ngay_sinh': parse_date(
                get_val(row, 'ngay_sinh') or get_val(row, 'Ngày sinh')
            ),
            'gioi_tinh': get_val(row, 'gioi_tinh') or get_val(row, 'Giới tính'),
            'dan_toc': get_val(row, 'dan_toc') or get_val(row, 'Dân tộc'),
            'email_sv': get_val(row, 'email_sv') or get_val(row, 'Email SV'),
            # Sửa lỗi: Lấy đúng cột điện thoại thay vì lặp lại email_sv
            'dien_thoai': (
                get_val(row, 'dien_thoai')
                or get_val(row, 'DienThoai')
                or get_val(row, 'Điện thoại')
                or get_val(row, 'Số điện thoại')
                or get_val(row, 'SDT')
            ),
            'mssv': get_val(row, 'mssv') or get_val(row, 'MSSV'),
            'mat_khau_online': get_val(row, 'mat_khau_online')
            or get_val(row, 'Mật khẩu Online'),
            'email_ueh': get_val(row, 'email_ueh') or get_val(row, 'Email UEH'),
            'mat_khau_email': get_val(row, 'mat_khau_email')
            or get_val(row, 'Mật khẩu Email'),
            'tong_diem': parse_float(
                get_val(row, 'tong_diem') or get_val(row, 'Tổng điểm')
            ),
            'dtc0_pt1': parse_float(
                get_val(row, 'dtc0_pt1') or get_val(row, 'ĐTC0 PT1')
            ),
            'dtc0_pt2': parse_float(
                get_val(row, 'dtc0_pt2') or get_val(row, 'ĐTC0 PT2')
            ),
            'loai_ccqt_uid': (
                get_val(row, 'loai_ccqt_uid')
                or get_val(row, 'LoaiCCQTUID')
                or get_val(row, 'Loại CCQT UID')
            ),
            'so_diem_ccqt': parse_float(
                get_val(row, 'so_diem_ccqt') or get_val(row, 'Số điểm CCQT')
            ),
            'ngay_thi_ccqt': parse_date(
                get_val(row, 'ngay_thi_ccqt') or get_val(row, 'Ngày thi CCQT')
            ),
            'ngay_thi_ccqt_new': parse_date(
                get_val(row, 'ngay_thi_ccqt_new')
                or get_val(row, 'Ngày thi CCQT Mới')
            ),
            'diem_tb_lop10': parse_float(
                get_val(row, 'diem_tb_lop10') or get_val(row, 'ĐTB Lớp 10')
            ),
            'diem_tb_lop11': parse_float(
                get_val(row, 'diem_tb_lop11') or get_val(row, 'ĐTB Lớp 11')
            ),
            'diem_tb_lop12': parse_float(
                get_val(row, 'diem_tb_lop12') or get_val(row, 'ĐTB Lớp 12')
            ),
        }

        if hasattr(CapNhatThongTinTrungTuyen, 'ho_ten'):
          kwargs['ho_ten'] = (
              get_val(row, 'ho_ten')
              or get_val(row, 'Họ và tên')
              or get_val(row, 'Họ tên')
          )

        dict_trung_tuyen[cccd] = CapNhatThongTinTrungTuyen(**kwargs)

      danh_sach_trung_tuyen = list(dict_trung_tuyen.values())

      if danh_sach_trung_tuyen:
        with transaction.atomic():
          CapNhatThongTinTrungTuyen.objects.all().delete()
          CapNhatThongTinTrungTuyen.objects.bulk_create(
              danh_sach_trung_tuyen, batch_size=1000
          )
        messages.success(
            request,
            f'Đã import thành công {len(danh_sach_trung_tuyen)} dòng dữ liệu!',
        )
      else:
        messages.warning(
            request, 'Không tìm thấy dữ liệu hợp lệ trong file Excel!'
        )

    except Exception as e:
      messages.error(request, f'Lỗi import file Excel: {e}')
    return redirect('danh_sach_cap_nhat_trung_tuyen')

  # Processing DataTables AJAX request
  if (
      request.headers.get('x-requested-with') == 'XMLHttpRequest'
      or request.GET.get('draw')
  ):
    draw = int(request.GET.get('draw', 1))
    start = int(request.GET.get('start', 0))
    length = int(request.GET.get('length', 10))
    search_val = request.GET.get('search[value]', '').strip()

    queryset = CapNhatThongTinTrungTuyen.objects.all()
    records_total = queryset.count()

    if search_val:
      filter_q = (
          Q(cccd__icontains=search_val)
          | Q(ma_dkxt__icontains=search_val)
          | Q(mssv__icontains=search_val)
          | Q(email_sv__icontains=search_val)
          | Q(sbd__icontains=search_val)
          | Q(ma_noi_sinh__icontains=search_val)
      )
      if hasattr(CapNhatThongTinTrungTuyen, 'ho_ten'):
        filter_q |= Q(ho_ten__icontains=search_val)
      queryset = queryset.filter(filter_q)

    records_filtered = queryset.count()
    data_slice = queryset[start : start + length] if length != -1 else queryset

    data = []
    for idx, item in enumerate(data_slice, start=start + 1):
      data.append({
          'id': item.id,
          'stt': idx,
          'cccd': getattr(item, 'cccd', '') or '',
          'ho_ten': (
              getattr(item, 'ho_ten', '')
              if hasattr(item, 'ho_ten')
              else ''
          ),
          'ma_dkxt': item.ma_dkxt or '',
          'sbd': getattr(item, 'sbd', '') or '',
          'ma_noi_sinh': getattr(item, 'ma_noi_sinh', '') or '',
          'ten_giai_hsg': getattr(item, 'ten_giai_hsg', '') or '',
          'ten_mon_hsg': getattr(item, 'ten_mon_hsg', '') or '',
          'ten_hang_hsg': getattr(item, 'ten_hang_hsg', '') or '',
          'nam_hsg': getattr(item, 'nam_hsg', '') or '',
          'ngay_sinh': (
              item.ngay_sinh.strftime('%d/%m/%Y') if item.ngay_sinh else ''
          ),
          'gioi_tinh': item.gioi_tinh or '',
          'dan_toc': item.dan_toc or '',
          'email_sv': item.email_sv or '',
          'dien_thoai': item.dien_thoai or '',
          'mssv': item.mssv or '',
          'mat_khau_online': item.mat_khau_online or '',
          'email_ueh': item.email_ueh or '',
          'mat_khau_email': item.mat_khau_email or '',
          'tong_diem': (
              item.tong_diem if item.tong_diem is not None else ''
          ),
          'dtc0_pt1': item.dtc0_pt1 if item.dtc0_pt1 is not None else '',
          'dtc0_pt2': item.dtc0_pt2 if item.dtc0_pt2 is not None else '',
          'loai_ccqt_uid': (
              getattr(item, 'loai_ccqt_uid', '')
              if hasattr(item, 'loai_ccqt_uid')
              else ''
          ),
          'so_diem_ccqt': (
              item.so_diem_ccqt if item.so_diem_ccqt is not None else ''
          ),
          'ngay_thi_ccqt': (
              item.ngay_thi_ccqt.strftime('%d/%m/%Y')
              if item.ngay_thi_ccqt
              else ''
          ),
          'ngay_thi_ccqt_new': (
              item.ngay_thi_ccqt_new.strftime('%d/%m/%Y')
              if item.ngay_thi_ccqt_new
              else ''
          ),
          'diem_tb_lop10': (
              item.diem_tb_lop10 if item.diem_tb_lop10 is not None else ''
          ),
          'diem_tb_lop11': (
              item.diem_tb_lop11 if item.diem_tb_lop11 is not None else ''
          ),
          'diem_tb_lop12': (
              item.diem_tb_lop12 if item.diem_tb_lop12 is not None else ''
          ),
      })

    return JsonResponse({
        'draw': draw,
        'recordsTotal': records_total,
        'recordsFiltered': records_filtered,
        'data': data,
    })

  danh_sach = CapNhatThongTinTrungTuyen.objects.all().order_by('id')
# Tính toán số lượng thực tế từ cơ sở dữ liệu
  stats = {
       'total': danh_sach.count(),
       'chua_co_mssv': danh_sach.filter(Q(mssv__isnull=True) | Q(mssv='')).count(),
       'chua_co_email': danh_sach.filter(Q(email_ueh__isnull=True) | Q(email_ueh='')).count(),
       'chua_co_ngay_sinh': danh_sach.filter(ngay_sinh__isnull=True).count(),
       'chua_co_noi_sinh': danh_sach.filter(Q(ma_noi_sinh__isnull=True) | Q(ma_noi_sinh='')).count(),
   }  
  return render(request, 'xettuyen/cap_nhat_thong_tin_trung_tuyen.html', {
        'danh_sach': danh_sach,
        'stats': stats,
    })


# =========================================================
# 3. XUẤT EXCEL DANH SÁCH CẬP NHẬT TRÚNG TUYỂN
# =========================================================

@custom_login_required
@check_permission('cap_nhat_trung_tuyen')
def export_cap_nhat_trung_tuyen(request):
    """Xuất danh sách cập nhật trúng tuyển ra Excel"""
    # Khởi tạo response tự động gắn ngày: "Danh_sach_cap_nhat_trung_tuyen_DDMMYYYY.xlsx"
    response = create_file_response('Danh_sach_cap_nhat_trung_tuyen', extension='xlsx')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CapNhatTrungTuyen"

    ws.append([
        'STT', 'CCCD', 'Mã ĐKXT', 'SBD', 'Mã nơi sinh', 'Tên giải HSG', 'Tên môn HSG', 'Tên hạng HSG', 'Năm HSG',
        'Ngày sinh', 'Giới tính', 'Dân tộc', 'Email SV', 'Điện thoại', 'MSSV',
        'Mật khẩu Online', 'Email UEH', 'Mật khẩu Email', 'Tổng điểm', 'ĐTC0 PT1', 'ĐTC0 PT2',
        'Loại CCQT UID', 'Số điểm CCQT', 'Ngày thi CCQT', 'Ngày thi CCQT Mới',
        'ĐTB Lớp 10', 'ĐTB Lớp 11', 'ĐTB Lớp 12'
    ])

    def fmt_date(val):
        if not val:
            return ''
        if hasattr(val, 'strftime'):
            return val.strftime('%d/%m/%Y')
        return str(val)

    queryset = CapNhatThongTinTrungTuyen.objects.all().order_by('-id')
    for idx, item in enumerate(queryset, 1):
        ws.append([
            idx,
            getattr(item, 'cccd', '') or '',
            getattr(item, 'ma_dkxt', '') or '',
            getattr(item, 'sbd', '') or '',
            getattr(item, 'ma_noi_sinh', '') or '',
            getattr(item, 'ten_giai_hsg', '') or '',
            getattr(item, 'ten_mon_hsg', '') or '',
            getattr(item, 'ten_hang_hsg', '') or '',
            getattr(item, 'nam_hsg', '') or '',
            fmt_date(item.ngay_sinh),
            getattr(item, 'gioi_tinh', '') or '',
            getattr(item, 'dan_toc', '') or '',
            getattr(item, 'email_sv', '') or '',
            getattr(item, 'dien_thoai', '') or '',
            getattr(item, 'mssv', '') or '',
            getattr(item, 'mat_khau_online', '') or '',
            getattr(item, 'email_ueh', '') or '',
            getattr(item, 'mat_khau_email', '') or '',
            item.tong_diem if item.tong_diem is not None else '',
            item.dtc0_pt1 if item.dtc0_pt1 is not None else '',
            item.dtc0_pt2 if item.dtc0_pt2 is not None else '',
            item.loai_ccqt_uid if item.loai_ccqt_uid is not None else '',
            item.so_diem_ccqt if item.so_diem_ccqt is not None else '',
            fmt_date(item.ngay_thi_ccqt),
            fmt_date(item.ngay_thi_ccqt_new),
            item.diem_tb_lop10 if item.diem_tb_lop10 is not None else '',
            item.diem_tb_lop11 if item.diem_tb_lop11 is not None else '',
            item.diem_tb_lop12 if item.diem_tb_lop12 is not None else '',
        ])

    wb.save(response)
    return response

# =========================================================
# 4. SỬA THÔNG TIN CẬP NHẬT TRÚNG TUYỂN
# =========================================================
@custom_login_required
@check_permission('cap_nhat_trung_tuyen')
def sua_cap_nhat_trung_tuyen(request, pk):
  """Sửa thông tin cập nhật trúng tuyển bằng Django ORM an toàn."""
  if request.method == 'POST':
    try:
      item = CapNhatThongTinTrungTuyen.objects.get(pk=pk)

      # 1. Thông tin cá nhân & Định danh
      item.cccd = request.POST.get('cccd', '').strip()
      if hasattr(item, 'ho_ten'):
        item.ho_ten = request.POST.get('ho_ten', '').strip()
      item.ma_dkxt = request.POST.get('ma_dkxt', '').strip()
      item.sbd = request.POST.get('sbd', '').strip() or None
      item.ngay_sinh = parse_date(request.POST.get('ngay_sinh'))
      item.gioi_tinh = request.POST.get('gioi_tinh', '').strip() or None
      item.dan_toc = request.POST.get('dan_toc', '').strip() or None
      item.ma_noi_sinh = request.POST.get('ma_noi_sinh', '').strip() or None

      # 2. Giải HSG
      item.ten_giai_hsg = request.POST.get('ten_giai_hsg', '').strip() or None
      item.ten_mon_hsg = request.POST.get('ten_mon_hsg', '').strip() or None
      item.ten_hang_hsg = request.POST.get('ten_hang_hsg', '').strip() or None
      item.nam_hsg = request.POST.get('nam_hsg', '').strip() or None

      # 3. Tài khoản, Email & Điện thoại
      item.email_sv = request.POST.get('email_sv', '').strip() or None
      item.dien_thoai = request.POST.get('dien_thoai', '').strip() or None
      item.mssv = request.POST.get('mssv', '').strip() or None
      item.mat_khau_online = (
          request.POST.get('mat_khau_online', '').strip() or None
      )
      item.email_ueh = request.POST.get('email_ueh', '').strip() or None
      item.mat_khau_email = (
          request.POST.get('mat_khau_email', '').strip() or None
      )

      # 4. Điểm số (ép kiểu float, tự chuyển dấu ',' thành '.')
      item.tong_diem = parse_float(request.POST.get('tong_diem'))
      item.dtc0_pt1 = parse_float(request.POST.get('dtc0_pt1'))
      item.dtc0_pt2 = parse_float(request.POST.get('dtc0_pt2'))
      item.diem_tb_lop10 = parse_float(request.POST.get('diem_tb_lop10'))
      item.diem_tb_lop11 = parse_float(request.POST.get('diem_tb_lop11'))
      item.diem_tb_lop12 = parse_float(request.POST.get('diem_tb_lop12'))

      # 5. Chứng chỉ quốc tế
      item.loai_ccqt_uid = request.POST.get('loai_ccqt_uid', '').strip() or None
      item.so_diem_ccqt = parse_float(request.POST.get('so_diem_ccqt'))
      item.ngay_thi_ccqt = parse_date(request.POST.get('ngay_thi_ccqt'))
      item.ngay_thi_ccqt_new = parse_date(
          request.POST.get('ngay_thi_ccqt_new')
      )

      item.save()
      messages.success(
          request, f'Cập nhật thông tin thí sinh {item.cccd} thành công!'
      )

    except CapNhatThongTinTrungTuyen.DoesNotExist:
      messages.error(request, 'Không tìm thấy dữ liệu thí sinh cần cập nhật!')
    except Exception as e:
      messages.error(request, f'Lỗi cập nhật dữ liệu: {e}')

  return redirect('danh_sach_cap_nhat_trung_tuyen')


# =========================================================
# 5. XÓA HỒ SƠ TRÚNG TUYỂN
# =========================================================

@custom_login_required
@check_permission('cap_nhat_trung_tuyen')
def xoa_cap_nhat_trung_tuyen(request, pk):
    """Xóa hồ sơ cập nhật trúng tuyển bằng request.POST"""
    if request.method == 'POST':
        with connection.cursor() as cursor:
            cursor.execute("DELETE FROM cap_nhat_thong_tin_trung_tuyen WHERE id = %s", [pk])

        messages.success(request, "Đã xóa hồ sơ trúng tuyển!")
    return redirect('danh_sach_cap_nhat_trung_tuyen')
 
@custom_login_required
@check_permission('import_bo_sung_trung_tuyen')
def import_bo_sung_trung_tuyen(request):
    """Đọc file Excel và lưu/cập nhật dữ liệu vào bảng mau_import_giay_bao theo CCCD"""
    if request.method == 'POST' and request.FILES.get('file_excel'):
        excel_file = request.FILES['file_excel']

        try:
            df = pd.read_excel(excel_file, dtype=str)

            if df.empty:
                messages.warning(request, 'File Excel không có dữ liệu.')
                return redirect(request.META.get('HTTP_REFERER', '/'))

            rename_dict = {}
            for col in df.columns:
                clean_col = str(col).strip().lower()
                if clean_col in EXCEL_TO_MAU_GIAY_BAO:
                    rename_dict[col] = EXCEL_TO_MAU_GIAY_BAO[clean_col]

            df = df.rename(columns=rename_dict)

            valid_columns = [c for c in df.columns if c in EXCEL_TO_MAU_GIAY_BAO.values()]
            if 'cccd' not in valid_columns:
                messages.error(
                    request,
                    'File Excel bắt buộc phải có cột `CCCD` để làm khóa định danh!',
                )
                return redirect(request.META.get('HTTP_REFERER', '/'))

            df_valid = df[valid_columns]

            cols_formatted = ', '.join([f'`{col}`' for col in valid_columns])
            placeholders = ', '.join(['%s'] * len(valid_columns))
            update_clause = ', '.join([
                f'`{col}` = VALUES(`{col}`)'
                for col in valid_columns
                if col != 'cccd'
            ])

            sql = f"""
                INSERT INTO `mau_import_giay_bao` ({cols_formatted})
                VALUES ({placeholders})
                ON DUPLICATE KEY UPDATE {update_clause}
            """

            values = [
                tuple(
                    None
                    if pd.isna(val) or str(val).strip() == ''
                    else str(val).strip()
                    for val in row
                )
                for row in df_valid.to_numpy()
            ]

            with connection.cursor() as cursor:
                cursor.executemany(sql, values)

            messages.success(
                request,
                f'Đã cập nhật dữ liệu thành công {len(df_valid)} bản ghi vào bảng `mau_import_giay_bao`!',
            )

        except Exception as e:
            messages.error(request, f'Lỗi khi import giấy báo: {str(e)}')

    return redirect(request.META.get('HTTP_REFERER', '/'))


# 4. XUẤT EXCEL DỮ LIỆU BẢNG MAU_IMPORT_GIAY_BAO
@custom_login_required
@check_permission('export_trung_tuyen')
def xuat_excel_trung_tuyen(request):
    # Khởi tạo response tự động gắn ngày: "Danh_Sach_Trung_Tuyen_DDMMYYYY.xlsx"
    response = create_file_response('Danh_Sach_Trung_Tuyen', extension='xlsx')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DS Trung Tuyen"

    headers = [
        "STT", "IDSV", "Số CV", "Họ tên", "Email", "Số điện thoại", "Mã ĐKXT", "Ngày sinh",
        "ĐTUT", "KVUT", "Đối tượng xét", "PTXT", "CTĐT", "Điểm cộng", "Học bạ", "THM TN",
        "Môn 1 (TN)", "Điểm 1", "Môn 2 (TN)", "Điểm 2", "Môn 3 (TN)", "Điểm 3",
        "Điểm ĐGNL", "THM VSAT", "Môn 1 (VSAT)", "Điểm 1", "Môn 2 (VSAT)", "Điểm 2",
        "Môn 3 (VSAT)", "Điểm 3", "Điểm TB THPT", "Điểm QĐ 1", "Điểm QĐ 2",
        "Điểm Cộng", "Điểm Chuẩn", "CCCD", "Mã SV", "Barcode", "Page"
    ]
    
    ws.append(headers)
    
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    danh_sach = MauImportGiayBao.objects.all().order_by('id')
    for idx, ts in enumerate(danh_sach, start=1):
        sdt_val = getattr(ts, 'dien_thoai', '') or getattr(ts, 'sdt', '') or ''
        ws.append([
            idx, ts.IDSV or '', ts.so_cv or '', ts.ho_ten or '', 
            ts.email or '', sdt_val, ts.ma_dkxt or '', ts.ngay_sinh or '', ts.dtut or '',
            ts.kvut or '', ts.phuong_thuc_xet or '', ts.ptxt or '', ts.ctdt or '', ts.pt2_diem_cong or '',
            ts.hoc_ba or '', ts.pt2_tn_thm or '', ts.pt2_tn_mamon1 or '',
            ts.pt2_tn_diemmon1 or '', ts.pt2_tn_mamon2 or '', ts.pt2_tn_diemmon2 or '',
            ts.pt2_tn_mamon3 or '', ts.pt2_tn_diemmon3 or '', ts.pt2_dgnl or '',
            ts.pt2_vsat_thm or '', ts.pt2_vsat_mamon1 or '', ts.pt2_vsat_diemmon1 or '',
            ts.pt2_vsat_mamon2 or '', ts.pt2_vsat_diemmon2 or '', ts.pt2_vsat_mamon3 or '',
            ts.pt2_vsat_diemmon3 or '', ts.pt2a_diemtbthpt or '', ts.pt2_diem_qd or '',
            ts.pt2_qd or '', ts.dtc0 or 0, ts.dc or '', ts.cccd or '', ts.ma_sv or '', ts.barcode or '', ts.page or ''
        ])

    wb.save(response)
    return response
    
# XUẤT DỮ LIỆU SMS
def bo_dau_tieng_viet(text):
    """Xóa hoàn toàn dấu tiếng Việt (kể cả chữ HOA và Đ/đ)"""
    if not text:
        return ""
    text = str(text).replace('Đ', 'D').replace('đ', 'd')
    text = unicodedata.normalize('NFD', text)
    # Phải dùng 'Mn' (chữ n viết thường) để lọc chính xác ký tự dấu
    text = ''.join(c for c in text if unicodedata.category(c) != 'Mn')
    return unicodedata.normalize('NFC', text).strip()


@custom_login_required
@check_permission('ds_trung_tuyen')
def export_sms(request):
    """Xuất file Excel gồm 3 cột: Họ tên không dấu, Số điện thoại, Nội dung SMS"""
    
    # Tạo sẵn HttpResponse chứa tên file "Danh_sach_xuat_SMS_11092026.xlsx"
    response = create_file_response('Danh_sach_xuat_SMS', extension='xlsx')

    # Lấy nội dung từ ô nhập liệu giao diện
    noi_dung_sms = request.POST.get('noi_dung_sms') or request.GET.get('noi_dung_sms') or ''

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DS Xuat SMS"

    # Định dạng 3 cột chính
    headers = ["Họ tên", "Điện thoại", "Nội dung"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    danh_sach = MauImportGiayBao.objects.all().order_by('id')
    for ts in danh_sach:
        sdt = str(getattr(ts, 'dien_thoai', '') or getattr(ts, 'sdt', '') or '').strip()
        ho_ten_khong_dau = bo_dau_tieng_viet(ts.ho_ten)

        ws.append([ho_ten_khong_dau, sdt, noi_dung_sms])

    wb.save(response)
    return response
# ==========================================
# HÀM THAY THẾ PLACEHOLDER TRONG FILE DOCX
# ==========================================
def replace_docx_placeholders(doc, context):
    bold_keys = ['HoTen', 'SoCV', 'NgaySinh', 'DT', 'Khuvuc', 'DienThoai']

    for p in doc.paragraphs:
        for key, value in context.items():
            placeholder = f'{{{key}}}'
            if placeholder in p.text:
                for run in p.runs:
                    if placeholder in run.text:
                        run.text = run.text.replace(placeholder, str(value))
                        if key in bold_keys:
                            run.bold = True

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for key, value in context.items():
                        placeholder = f'{{{key}}}'
                        if placeholder in p.text:
                            for run in p.runs:
                                if placeholder in run.text:
                                    run.text = run.text.replace(placeholder, str(value))
                                    if key in bold_keys:
                                        run.bold = True

    clean_context = {
        str(k).strip(): (str(v) if v is not None else '') for k, v in context.items()
    }

    def replace_in_xml_element(element):
        if element is None:
            return

        for p in element.xpath('.//*[local-name()="p"]'):
            if any(
                ancestor.tag.endswith(('Fallback', 'pict'))
                for ancestor in p.iterancestors()
            ):
                continue

            t_nodes = [
                t
                for t in p.xpath('.//*[local-name()="t"]')
                if next((a for a in t.iterancestors() if a.tag.endswith('p')), None) == p
            ]

            if not t_nodes:
                continue

            full_text = ''.join(node.text for node in t_nodes if node.text)
            full_text = full_text.replace('\xa0', ' ')

            if '{' in full_text and '}' in full_text:
                best_rPr = None
                for t in t_nodes:
                    r = t.getparent()
                    rPr = r.find(
                        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr'
                    )
                    if rPr is not None:
                        rPr_str = etree.tostring(rPr).decode('utf-8')
                        if any(
                            k in rPr_str
                            for k in [
                                'IDAutomation',
                                'Barcode',
                                'Code39',
                                'Code 39',
                                'HC39',
                            ]
                        ):
                            best_rPr = rPr
                            break

                for key, val in clean_context.items():
                    pattern = re.compile(
                        r'\{\s*' + re.escape(key) + r'\s*\}', re.IGNORECASE
                    )
                    full_text = pattern.sub(val, full_text)

                if best_rPr is not None:
                    first_r = t_nodes[0].getparent()
                    existing_rPr = first_r.find(
                        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr'
                    )
                    if existing_rPr is not None:
                        first_r.remove(existing_rPr)
                    first_r.insert(0, copy.deepcopy(best_rPr))

                t_nodes[0].text = full_text
                for node in t_nodes[1:]:
                    node.text = ''

    replace_in_xml_element(doc._element)

    for section in doc.sections:
        for hf in [
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ]:
            if hf and hasattr(hf, '_element'):
                replace_in_xml_element(hf._element)


# ==========================================
# HÀM CHUYỂN ĐỔI 1 FILE DOCX SANG PDF BYTES (TỰ ĐỘNG CHỌN MS WORD / LIBREOFFICE)
# ==========================================
def convert_single_docx_to_pdf_bytes(doc):
    """Lưu tạm Document và biên dịch sang PDF Bytes an toàn trên Windows/Linux"""
    temp_dir = os.path.join(settings.BASE_DIR, 'media', 'temp_single')
    os.makedirs(temp_dir, exist_ok=True)

    unique_id = uuid.uuid4().hex
    docx_path = os.path.abspath(os.path.join(temp_dir, f'temp_{unique_id}.docx'))
    pdf_path = os.path.abspath(os.path.join(temp_dir, f'temp_{unique_id}.pdf'))

    # Lưu file DOCX tạm
    doc.save(docx_path)

    profile_dir = f"/tmp/libreoffice_profile_{os.getpid()}_{unique_id}"

    try:
        if sys.platform == 'win32':
            import pythoncom
            import win32com.client

            pythoncom.CoInitialize()
            word = None
            try:
                word = win32com.client.DispatchEx('Word.Application')
                word.Visible = False
                word.DisplayAlerts = 0

                word_doc = word.Documents.Open(
                    FileName=docx_path,
                    ConfirmConversions=False,
                    ReadOnly=True,
                    AddToRecentFiles=False,
                )
                word_doc.SaveAs(pdf_path, FileFormat=17)  # 17 = wdFormatPDF
                word_doc.Close(SaveChanges=0)
            except Exception as e:
                raise Exception(f'Lỗi khi xuất PDF từ MS Word: {str(e)}')
            finally:
                if word:
                    try:
                        word.Quit()
                    except Exception:
                        pass
                pythoncom.CoUninitialize()
        else:
            # Môi trường Linux Server (Render, Docker, VPS)
            cmd = [
                'libreoffice',
                f'-env:UserInstallation=file://{profile_dir}',
                '--headless',
                '--convert-to',
                'pdf',
                docx_path,
                '--outdir',
                temp_dir,
            ]
            result = subprocess.run(
                cmd,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                timeout=60,
            )

            if result.returncode != 0:
                raise RuntimeError(f"Lỗi LibreOffice: {result.stderr or result.stdout}")

        if not os.path.exists(pdf_path):
            raise FileNotFoundError(f'Không thể tạo file PDF tại {pdf_path}')

        with open(pdf_path, 'rb') as f:
            pdf_bytes = f.read()

        return pdf_bytes

    finally:
        # Dọn dẹp các file tạm
        for path in [docx_path, pdf_path]:
            if os.path.exists(path):
                try:
                    os.remove(path)
                except Exception:
                    pass
        if sys.platform != 'win32' and os.path.exists(profile_dir):
            try:
                shutil.rmtree(profile_dir, ignore_errors=True)
            except Exception:
                pass


# ==========================================
# PREPARE DOCX DOCUMENT
# ==========================================
def prepare_docx_document(ts):
    """Tạo đối tượng Document Word đã được điền dữ liệu thí sinh"""
    ptxt_upper = (ts.ptxt or '').upper()

    if 'DGNL' in ptxt_upper or (ts.pt2_dgnl and str(ts.pt2_dgnl).strip() != ''):
        template_name = 'DS PT2-DGNL.docx'
    elif 'VSAT' in ptxt_upper or (ts.pt2_vsat_thm and str(ts.pt2_vsat_thm).strip() != ''):
        template_name = 'DS PT2-VSAT.docx'
    else:
        template_name = 'DS PT2-TN THPT.docx'

    template_path = os.path.join(
        settings.BASE_DIR, 'xettuyen', 'templates', 'docx', template_name
    )

    raw_ctdt = str(ts.ctdt or '').strip()
    ctdt_ten = raw_ctdt.split(':', 1)[-1].strip() if raw_ctdt else ''

    raw_barcode = (
        str(ts.barcode).strip()
        if getattr(ts, 'barcode', None)
        else (ts.cccd or ts.ma_dkxt or '')
    )
    formatted_barcode = f'*{raw_barcode.strip("*")}*' if raw_barcode else ''

    context = {
        'SoCV': ts.so_cv or '',
        'HoTen': ts.ho_ten or '',
        'MaDKXT': ts.ma_dkxt or '',
        'NgaySinh': ts.ngay_sinh or '',
        'DienThoai': getattr(ts, 'dien_thoai', '') or getattr(ts, 'sdt', '') or '',
        'DT': ts.dtut or '',
        'Khuvuc': ts.kvut or '',
        'HocBa': ts.hoc_ba or '',
        'phuong_thuc_xet': ts.phuong_thuc_xet or '',
        'CTDT': ctdt_ten,
        'CCQT': getattr(ts, 'ccqt', '') or '',
        'PT2_DiemCong': getattr(ts, 'pt2_diem_cong', '') or '',
        'PT2_TN_THM': ts.pt2_tn_thm or '',
        'PT2_TN_MaMon1': ts.pt2_tn_mamon1 or '',
        'PT2_TN_DiemMon1': ts.pt2_tn_diemmon1 or '',
        'PT2_TN_MaMon2': ts.pt2_tn_mamon2 or '',
        'PT2_TN_DiemMon2': ts.pt2_tn_diemmon2 or '',
        'PT2_TN_MaMon3': ts.pt2_tn_mamon3 or '',
        'PT2_TN_DiemMon3': ts.pt2_tn_diemmon3 or '',
        'PT2_VSAT_THM': ts.pt2_vsat_thm or '',
        'PT2_VSAT_MaMon1': ts.pt2_vsat_mamon1 or '',
        'PT2_VSAT_DiemMon1': ts.pt2_vsat_diemmon1 or '',
        'PT2_VSAT_MaMon2': ts.pt2_vsat_mamon2 or '',
        'PT2_VSAT_DiemMon2': ts.pt2_vsat_diemmon2 or '',
        'PT2_VSAT_MaMon3': ts.pt2_vsat_mamon3 or '',
        'PT2_VSAT_DiemMon3': ts.pt2_vsat_diemmon3 or '',
        'PT2_DGNL': ts.pt2_dgnl or '',
        'PT2a_DiemTBTHPT': ts.pt2a_diemtbthpt or '',
        'PT2_DiemQD': ts.pt2_diem_qd or '',
        'PT2_QD': ts.pt2_qd or '',
        'BarCode': formatted_barcode,
        'DTC0': ts.dtc0 if ts.dtc0 is not None else '0',
        'DC': ts.dc if ts.dc is not None else '',
        'Page': ts.page or '',
        'ViTri_1': '',
        'ViTri_2': '',
    }

    doc = Document(template_path)
    replace_docx_placeholders(doc, context)
    return doc


# ==========================================
# VIEW IN GIẤY BÁO TRÚNG TUYỂN
# ==========================================
@custom_login_required
@check_permission('in_giay_bao')
def xu_ly_xuat_giay_bao(request, cccd=None, ts_id=None):
    identifier = cccd or ts_id

    # 1. Kiểm tra quyền truy cập (Đã bổ sung role_code 'canbo')
    role_code = (request.session.get('role_code') or '').lower()
    role_name = request.session.get('role_name')

    is_staff_or_admin = (
        role_code in ['admin', 'canbo']
        or role_name in ['Quản trị viên', 'Administrator', 'Cán bộ', 'Can bo']
    )
    student_cccd = request.session.get('student_cccd')

    # Chỉ chặn nếu KHÔNG PHẢI cán bộ/admin VÀ (không phải thí sinh hoặc cố xem CCCD khác)
    if not is_staff_or_admin and (
        not student_cccd or str(student_cccd) != str(identifier)
    ):
        return HttpResponse('Bạn không có quyền truy cập thông tin này!', status=403)

    is_preview = request.GET.get('preview') == '1'
    fetch_pdf = request.GET.get('fetch_pdf') == '1'

    # ⚡ BƯỚC 1: NẾU BẤM XEM TRƯỚC -> LUÔN HIỂN THỊ MÀN HÌNH LOADING TRƯỚC
    if is_preview and not fetch_pdf:
        loading_html = f"""
        <!DOCTYPE html>
        <html lang="vi">
        <head>
            <meta charset="UTF-8">
            <meta name="viewport" content="width=device-width, initial-scale=1.0">
            <title>Đang tải Giấy báo - {identifier}</title>
            <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
            <link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.4.0/css/all.min.css">
            <style>
                body {{
                    background-color: #f4f6f9;
                    height: 100vh;
                    display: flex;
                    align-items: center;
                    justify-content: center;
                    font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
                }}
                .loading-card {{
                    background: #ffffff;
                    padding: 40px;
                    border-radius: 16px;
                    box-shadow: 0 10px 30px rgba(0, 0, 0, 0.08);
                    text-align: center;
                    max-width: 460px;
                    width: 90%;
                }}
                .spinner-border {{
                    width: 3.5rem;
                    height: 3.5rem;
                }}
            </style>
        </head>
        <body>
            <div class="loading-card">
                <div class="spinner-border text-primary mb-4" role="status"></div>
                <h5 class="fw-bold mb-2 text-dark">Đang nạp Giấy báo trúng tuyển</h5>
                <p class="text-muted small mb-3">Vui lòng chờ trong giây lát...</p>

                <div class="progress mb-3" style="height: 12px; border-radius: 6px;">
                    <div id="progressBar" class="progress-bar progress-bar-striped progress-bar-animated bg-primary" style="width: 30%;"></div>
                </div>

                <small id="statusText" class="text-secondary fw-semibold">Đang chuẩn bị dữ liệu...</small>
            </div>

            <script>
                const progressBar = document.getElementById('progressBar');
                const statusText = document.getElementById('statusText');

                let progress = 30;
                const interval = setInterval(() => {{
                    if (progress < 85) {{
                        progress += 15;
                        progressBar.style.width = progress + '%';
                    }}
                }}, 150);

                const requestUrl = new URL(window.location.href);
                requestUrl.searchParams.set('fetch_pdf', '1');

                fetch(requestUrl.toString())
                    .then(response => {{
                        if (!response.ok) throw new Error('Không thể tải file PDF.');
                        return response.blob();
                    }})
                    .then(blob => {{
                        clearInterval(interval);
                        progressBar.style.width = '100%';
                        statusText.innerText = 'Hoàn tất! Đang mở PDF...';

                        const pdfUrl = URL.createObjectURL(blob);
                        setTimeout(() => {{
                            window.location.href = pdfUrl;
                        }}, 150);
                    }})
                    .catch(err => {{
                        clearInterval(interval);
                        document.body.innerHTML = `
                            <div class="card p-4 shadow-sm border-0 text-center" style="max-width: 500px; margin: auto;">
                                <div class="text-danger mb-3"><i class="fa-solid fa-triangle-exclamation fa-3x"></i></div>
                                <h5 class="text-danger fw-bold">⚠️ Không thể tạo file PDF</h5>
                                <p class="text-secondary small mt-2">${{err.message}}</p>
                                <button class="btn btn-primary btn-sm mt-2" onclick="location.reload()">Thử lại</button>
                            </div>
                        `;
                    }});
            </script>
        </body>
        </html>
        """
        return HttpResponse(loading_html)

    # ⚡ BƯỚC 2: KIỂM TRA CACHE KHI JAVASCRIPT GỌI FETCH_PDF=1 (HOẶC TẢI WORD)
    cache_key = (
        f'giay_bao_pdf_{identifier}' if is_preview else f'giay_bao_docx_{identifier}'
    )
    cached_file = cache.get(cache_key)

    if cached_file:
        if is_preview:
            response = HttpResponse(cached_file, content_type='application/pdf')
            response['Content-Disposition'] = (
                f'inline; filename="XemTruoc_{identifier}.pdf"'
            )
            return response
        else:
            response = HttpResponse(
                cached_file,
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            )
            response['Content-Disposition'] = (
                f'attachment; filename="GiayBao_{identifier}.docx"'
            )
            return response

    # ⚡ BƯỚC 3: NẾU CHƯA CÓ CACHE THÌ MỚI TRUY VẤN DB & TẠO FILE TỪ MS WORD
    if cccd:
        ts = MauImportGiayBao.objects.filter(cccd=cccd).first()
    else:
        ts = MauImportGiayBao.objects.filter(Q(id=ts_id) | Q(IDSV=ts_id)).first()

    if not ts:
        return HttpResponse('Không tìm thấy dữ liệu thí sinh!', status=404)

    doc = prepare_docx_document(ts)

    # ⚡ BƯỚC 4: SINH FILE LẦN ĐẦU VÀ LƯU VÀO CACHE TRONG 4 GIỜ
    if is_preview:
        try:
            pdf_bytes = convert_single_docx_to_pdf_bytes(doc)
            cache.set(cache_key, pdf_bytes, timeout=14400)  # Lưu Cache 4h

            response = HttpResponse(pdf_bytes, content_type='application/pdf')
            response['Content-Disposition'] = (
                f'inline; filename="XemTruoc_{ts.cccd or ts.ma_sv}.pdf"'
            )
            return response
        except Exception as e:
            return HttpResponse(f'Lỗi tạo PDF: {str(e)}', status=500)

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()
    cache.set(cache_key, docx_bytes, timeout=14400)  # Lưu Cache 4h

    filename = f"GiayBao_{ts.cccd or ts.ma_sv or 'ThongBao'}.docx"
    response = HttpResponse(
        docx_bytes,
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# ==========================================
# XUẤT TOÀN BỘ GIẤY BÁO TRÚNG TUYỂN TẤT CẢ QUA PDF ZIP
# ==========================================
def clean_filename(filename):
    return re.sub(r'[\\/*?:"<>|]', '', filename)


def generate_single_docx(ts, idx, docx_dir):
    raw_text = f"{ts.ptxt or ''} {ts.phuong_thuc_xet or ''}".upper()
    clean_text = raw_text.replace('-', '').replace('_', '').replace(' ', '')

    if 'DGNL' in clean_text or 'ĐGNL' in clean_text:
        category = 'DGNL'
    elif 'VSAT' in clean_text:
        category = 'VSAT'
    else:
        category = 'THPT'

    cat_docx_dir = os.path.join(docx_dir, category)
    os.makedirs(cat_docx_dir, exist_ok=True)

    try:
        doc = prepare_docx_document(ts)
        ident = ts.cccd or ts.ma_sv or str(idx)
        filename = clean_filename(f'{idx:04d}_{ident}_{ts.ho_ten}.docx')
        filepath = os.path.join(cat_docx_dir, filename)
        doc.save(filepath)
        return filepath
    except Exception as e:
        print(f'Lỗi tạo file docx thí sinh {ts.ho_ten}: {e}')
        return None


def convert_folder_docx_to_pdf(
    docx_dir, pdf_dir, task_id=None, total_files=1, zip_file_path=''
):
    converted_count = 0

    if sys.platform == 'win32':
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        word = None
        try:
            word = win32com.client.DispatchEx('Word.Application')
            word.Visible = False
            word.DisplayAlerts = 0

            for root, dirs, files in os.walk(docx_dir):
                rel_path = os.path.relpath(root, docx_dir)
                target_pdf_dir = (
                    pdf_dir if rel_path == '.' else os.path.join(pdf_dir, rel_path)
                )
                os.makedirs(target_pdf_dir, exist_ok=True)

                for filename in files:
                    if filename.endswith('.docx') and not filename.startswith('~$'):
                        docx_path = os.path.abspath(os.path.join(root, filename))
                        pdf_filename = os.path.splitext(filename)[0] + '.pdf'
                        pdf_path = os.path.abspath(
                            os.path.join(target_pdf_dir, pdf_filename)
                        )

                        try:
                            doc = word.Documents.Open(
                                FileName=docx_path,
                                ConfirmConversions=False,
                                ReadOnly=True,
                                AddToRecentFiles=False,
                            )
                            doc.SaveAs(pdf_path, FileFormat=17)
                            doc.Close(SaveChanges=0)

                            converted_count += 1
                            if task_id and total_files > 0:
                                percent = 40 + int((converted_count / total_files) * 50)
                                cache.set(
                                    f'task_{task_id}',
                                    {
                                        'status': 'processing',
                                        'current': converted_count,
                                        'total': total_files,
                                        'percent': percent,
                                        'file_path': zip_file_path,
                                    },
                                    timeout=3600,
                                )
                        except Exception as doc_err:
                            print(f'Lỗi tạo PDF file {filename}: {doc_err}')
                            continue

            return True
        except Exception as e:
            raise Exception(f'Lỗi MS Word: {str(e)}')
        finally:
            if word:
                try:
                    word.Quit()
                except Exception:
                    pass
            pythoncom.CoUninitialize()
    else:
        # Môi trường Linux Server / Docker
        for root, dirs, files in os.walk(docx_dir):
            rel_path = os.path.relpath(root, docx_dir)
            target_pdf_dir = (
                pdf_dir if rel_path == '.' else os.path.join(pdf_dir, rel_path)
            )
            os.makedirs(target_pdf_dir, exist_ok=True)

            for filename in files:
                if filename.endswith('.docx') and not filename.startswith('~$'):
                    docx_path = os.path.abspath(os.path.join(root, filename))
                    unique_sub_id = uuid.uuid4().hex
                    profile_dir = f"/tmp/libreoffice_profile_{os.getpid()}_{unique_sub_id}"

                    cmd = [
                        'libreoffice',
                        f'-env:UserInstallation=file://{profile_dir}',
                        '--headless',
                        '--convert-to',
                        'pdf',
                        docx_path,
                        '--outdir',
                        target_pdf_dir,
                    ]
                    try:
                        subprocess.run(
                            cmd,
                            stdout=subprocess.PIPE,
                            stderr=subprocess.PIPE,
                            text=True,
                            timeout=60,
                        )
                        converted_count += 1
                        if task_id and total_files > 0:
                            percent = 40 + int((converted_count / total_files) * 50)
                            cache.set(
                                f'task_{task_id}',
                                {
                                    'status': 'processing',
                                    'current': converted_count,
                                    'total': total_files,
                                    'percent': percent,
                                    'file_path': zip_file_path,
                                },
                                timeout=3600,
                            )
                    except Exception as doc_err:
                        print(f'Lỗi tạo PDF file {filename}: {doc_err}')
                    finally:
                        if os.path.exists(profile_dir):
                            try:
                                shutil.rmtree(profile_dir, ignore_errors=True)
                            except Exception:
                                pass
        return True


def start_export_all_pdf(request):
    task_id = str(uuid.uuid4())
    ds_thi_sinh = list(MauImportGiayBao.objects.all())
    total = len(ds_thi_sinh)

    if total == 0:
        return JsonResponse(
            {'error': 'Không có dữ liệu thí sinh để xuất!'}, status=400
        )

    cache.set(
        f'task_{task_id}',
        {
            'status': 'processing',
            'current': 0,
            'total': total,
            'percent': 0,
            'file_path': '',
        },
        timeout=3600,
    )

    def run_batch():
        base_temp = os.path.join(settings.BASE_DIR, 'media', 'temp_export', task_id)
        docx_dir = os.path.join(base_temp, 'docx')
        pdf_dir = os.path.join(base_temp, 'pdf')
        zip_dir = os.path.join(settings.BASE_DIR, 'media', 'temp_zip')

        os.makedirs(docx_dir, exist_ok=True)
        os.makedirs(pdf_dir, exist_ok=True)
        os.makedirs(zip_dir, exist_ok=True)

        zip_file_path = os.path.join(zip_dir, f'GiayBao_{task_id}.zip')

        try:
            for idx, ts in enumerate(ds_thi_sinh, start=1):
                generate_single_docx(ts, idx, docx_dir)
                percent = int((idx / total) * 40)
                cache.set(
                    f'task_{task_id}',
                    {
                        'status': 'processing',
                        'current': idx,
                        'total': total,
                        'percent': percent,
                        'file_path': zip_file_path,
                    },
                    timeout=3600,
                )

            convert_folder_docx_to_pdf(
                docx_dir,
                pdf_dir,
                task_id=task_id,
                total_files=total,
                zip_file_path=zip_file_path,
            )

            cache.set(
                f'task_{task_id}',
                {
                    'status': 'processing',
                    'current': total,
                    'total': total,
                    'percent': 92,
                    'file_path': zip_file_path,
                },
                timeout=3600,
            )

            with zipfile.ZipFile(
                zip_file_path, 'w', zipfile.ZIP_DEFLATED
            ) as zip_file:
                for root, dirs, files in os.walk(pdf_dir):
                    for file in files:
                        if file.endswith('.pdf'):
                            file_full_path = os.path.join(root, file)
                            arcname = os.path.relpath(file_full_path, pdf_dir)
                            zip_file.write(file_full_path, arcname=arcname)

            try:
                shutil.rmtree(base_temp)
            except Exception:
                pass

            cache.set(
                f'task_{task_id}',
                {
                    'status': 'completed',
                    'current': total,
                    'total': total,
                    'percent': 100,
                    'file_path': zip_file_path,
                },
                timeout=3600,
            )

        except Exception as err:
            cache.set(
                f'task_{task_id}',
                {
                    'status': 'error',
                    'message': str(err),
                    'current': 0,
                    'total': total,
                    'percent': 0,
                    'file_path': '',
                },
                timeout=3600,
            )

    threading.Thread(target=run_batch).start()
    return JsonResponse({'task_id': task_id})


@custom_login_required
@check_permission('in_giay_bao')
def check_export_status(request, task_id):
    data = cache.get(f'task_{task_id}')
    if not data:
        return JsonResponse({'error': 'Nhiệm vụ không tồn tại.'}, status=404)
    return JsonResponse(data)


@custom_login_required
@check_permission('in_giay_bao')
def download_export_zip(request, task_id):
    data = cache.get(f'task_{task_id}')
    if not data or not os.path.exists(data.get('file_path', '')):
        raise Http404('File nén không tồn tại hoặc đã hết hạn.')

    file_path = data['file_path']

    response = create_file_response('DanhSach_GiayBao_PDF', extension='zip')

    with open(file_path, 'rb') as f:
        response.content = f.read()

    try:
        os.remove(file_path)
    except Exception:
        pass

    return response


# ==========================================
# PHÂN HỆ IN GIẤY BÁO TRÚNG TUYỂN (DANH SÁCH)
# ==========================================
@custom_login_required
@check_permission('in_giay_bao')
def ds_in_giay_bao(request):
    """View hiển thị danh sách thí sinh trong bảng mau_import_giay_bao cho DataTables."""
    cau_hinh = CauHinhGiayBao.get_config()

    all_data = MauImportGiayBao.objects.all()
    stats = {
        'total': all_data.count(),
        'thpt': all_data.filter(ptxt__icontains='THPT').count(),
        'vsat': all_data.filter(
            Q(ptxt__icontains='VSAT') | Q(ptxt__icontains='V-SAT')
        ).count(),
        'dgnl': all_data.filter(ptxt__icontains='DGNL').count(),
    }

    danh_sach = all_data.order_by('id')

    context = {
        'cau_hinh': cau_hinh,
        'danh_sach': danh_sach,
        'page_obj': None,  # Giữ giá trị None để tránh lỗi VariableDoesNotExist trên HTML
        'stats': stats,
    }
    return render(request, 'xettuyen/ds_in_giay_bao.html', context)


# ==========================================
# CODE LƯU CẤU HÌNH
# ==========================================
@custom_login_required
@check_permission('in_giay_bao')
def luu_cau_hinh_giay_bao(request):
    """Xử lý lưu thời gian mở/đóng và bật/tắt hiển thị giấy báo từ Form"""
    if request.method == 'POST':
        ngay_bat_dau_raw = request.POST.get('ngay_bat_dau')
        ngay_ket_thuc_raw = request.POST.get('ngay_ket_thuc')
        cho_phep_xem_raw = request.POST.get('cho_phep_xem')

        cau_hinh = CauHinhGiayBao.get_config()

        try:
            if ngay_bat_dau_raw:
                dt_start = parse_datetime(ngay_bat_dau_raw)
                if dt_start and timezone.is_naive(dt_start):
                    dt_start = timezone.make_aware(
                        dt_start, timezone.get_current_timezone()
                    )
                cau_hinh.ngay_bat_dau = dt_start
            else:
                cau_hinh.ngay_bat_dau = None

            if ngay_ket_thuc_raw:
                dt_end = parse_datetime(ngay_ket_thuc_raw)
                if dt_end and timezone.is_naive(dt_end):
                    dt_end = timezone.make_aware(
                        dt_end, timezone.get_current_timezone()
                    )
                cau_hinh.ngay_ket_thuc = dt_end
            else:
                cau_hinh.ngay_ket_thuc = None

            cau_hinh.cho_phep_xem = cho_phep_xem_raw in ['on', '1', 'true', True]

            cau_hinh.save()
            messages.success(
                request,
                'Đã lưu cấu hình thời gian và trạng thái tra cứu thành công!',
            )
        except Exception as e:
            messages.error(request, f'Lỗi khi lưu cấu hình: {str(e)}')

    return redirect('ds_in_giay_bao')


# ==========================================
# 1. VIEW CHO ADMIN VÀ CÁN BỘ (BẮT BUỘC ĐĂNG NHẬP)
# ==========================================
@custom_login_required
@check_permission('in_giay_bao')
def in_giay_bao(request, cccd=None, ts_id=None):
    identifier = cccd or ts_id

    role_code = (request.session.get('role_code') or '').lower()
    role_name = request.session.get('role_name')

    is_staff_or_admin = (
        role_code in ['admin', 'canbo']
        or role_name in ['Quản trị viên', 'Administrator', 'Cán bộ']
    )

    student_cccd = request.session.get('student_cccd')

    if not is_staff_or_admin and (
        not student_cccd or str(student_cccd) != str(identifier)
    ):
        return HttpResponse('Bạn không có quyền truy cập thông tin này!', status=403)

    return xu_ly_xuat_giay_bao(request, identifier)


# ==========================================
# 2. VIEW CHO SINH VIÊN TRA CỨU (DÙNG SESSION TRA CỨU)
# ==========================================
def in_giay_bao_sinh_vien(request):
    """Đọc cấu hình thời gian trong DB và kiểm tra quyền tra cứu của Sinh viên"""
    cau_hinh = CauHinhGiayBao.get_config()
    now = timezone.now()

    if cau_hinh.ngay_bat_dau and now < cau_hinh.ngay_bat_dau:
        start_str = cau_hinh.ngay_bat_dau.strftime('%H:%M %d/%m/%Y')
        return HttpResponse(
            f'Cổng tra cứu giấy báo chưa mở. Thời gian bắt đầu: {start_str}',
            status=403,
        )

    if cau_hinh.ngay_ket_thuc and now > cau_hinh.ngay_ket_thuc:
        return HttpResponse(
            'Hệ thống đã đóng cổng tra cứu và in giấy báo trúng tuyển!',
            status=403,
        )

    student_cccd = request.session.get('student_cccd')
    if not student_cccd:
        return HttpResponse(
            'Phiên làm việc hết hạn hoặc bạn chưa đăng nhập!', status=403
        )

    return xu_ly_xuat_giay_bao(request, student_cccd)
# DANH SÁCH TRƯỜNG THPT  
@custom_login_required
@check_permission('danh_sach_truong_thpt')
def danh_sach_truong_thpt(request):
    if request.method == 'POST' and request.FILES.get('excel_file'):
        file_excel = request.FILES['excel_file']
        try:
            df = pd.read_excel(file_excel, dtype=str).fillna('')
            with connection.cursor() as cursor:
                TruongTHPT.objects.all().delete()

            danh_sach_moi = []
            for _, row in df.iterrows():
                ma_tinh = str(row.get('matinh', '')).strip().split('.')[0].zfill(2)
                ma_truong = str(row.get('matruong', '')).strip().split('.')[0].zfill(3)
                ma_phuong_xa = str(row.get('maphuongxa', '')).strip().split('.')[0].zfill(5)
                ma_truong_ghep = str(row.get('matruongghep', '')).strip().split('.')[0].zfill(5)
                ten_truong = str(row.get('tentruong', '')).strip()

                if not ma_truong_ghep or not ten_truong:
                    continue

                danh_sach_moi.append(TruongTHPT(
                    ma_truong_ghep=ma_truong_ghep,
                    ma_tinh=ma_tinh,
                    ten_tinh=str(row.get('tentinhtp', '')).strip(),
                    ma_phuong_xa=ma_phuong_xa,
                    ten_xa_phuong=str(row.get('tenxaphuong', '')).strip(),
                    ma_truong=ma_truong,
                    ten_truong=ten_truong,
                    dia_chi=str(row.get('diachi', '')).strip(),
                    khu_vuc=str(row.get('khuvuc', '')).strip(),
                ))

            if danh_sach_moi:
                TruongTHPT.objects.bulk_create(danh_sach_moi, batch_size=1000)
                messages.success(request, f"Đã import thành công {len(danh_sach_moi)} trường THPT mới!")
        except Exception as e:
            messages.error(request, f"Lỗi xử lý file Excel: {e}")
        return redirect('danh_sach_truong_thpt')

    if request.headers.get('x-requested-with') == 'XMLHttpRequest' or request.GET.get('draw'):
        draw = int(request.GET.get('draw', 1))
        start = int(request.GET.get('start', 0))
        length = int(request.GET.get('length', 10))
        search_value = request.GET.get('search[value]', '').strip()

        queryset = TruongTHPT.objects.all()
        records_total = queryset.count()

        if search_value:
            queryset = queryset.filter(
                Q(ma_truong_ghep__icontains=search_value) |
                Q(ten_truong__icontains=search_value) |
                Q(ten_tinh__icontains=search_value) |
                Q(ten_xa_phuong__icontains=search_value) |
                Q(ma_truong__icontains=search_value)
            )

        records_filtered = queryset.count()

        order_column_index = request.GET.get('order[0][column]', '0')
        order_dir = request.GET.get('order[0][dir]', 'asc')
        columns_map = {
            '0': 'id',
            '1': 'ma_truong_ghep',
            '2': 'ma_tinh',
            '3': 'ten_tinh',
            '4': 'ma_phuong_xa',
            '5': 'ten_xa_phuong',
            '6': 'ma_truong',
            '7': 'ten_truong',
            '8': 'dia_chi',
            '9': 'khu_vuc',
        }
        order_column = columns_map.get(order_column_index, 'id')
        if order_dir == 'desc':
            order_column = '-' + order_column

        queryset = queryset.order_by(order_column)

        if length != -1:
            data_slice = queryset[start:start + length]
        else:
            data_slice = queryset

        data = []
        for idx, item in enumerate(data_slice, start=start + 1):
            data.append({
                'stt': idx,
                'id': item.id,
                'ma_truong_ghep': item.ma_truong_ghep or '',
                'ma_tinh': item.ma_tinh or '',
                'ten_tinh': item.ten_tinh or '',
                'ma_phuong_xa': item.ma_phuong_xa or '',
                'ten_xa_phuong': item.ten_xa_phuong or '',
                'ma_truong': item.ma_truong or '',
                'ten_truong': item.ten_truong or '',
                'dia_chi': item.dia_chi or '',
                'khu_vuc': item.khu_vuc or '',
            })

        return JsonResponse({
            'draw': draw,
            'recordsTotal': records_total,
            'recordsFiltered': records_filtered,
            'data': data,
        })

    total_count = TruongTHPT.objects.count()
    return render(request, 'xettuyen/truong_thpt.html', {'total_count': total_count})


@custom_login_required
@check_permission('export_truong_thpt')
def xuat_excel_truong_thpt(request):
    try:
        # Sử dụng hàm create_file_response để tự động sinh tên file dạng "Danh_sach_truong_THPT_DDMMYYYY.xlsx"
        response = create_file_response('Danh_sach_truong_THPT', extension='xlsx')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Danh sách trường THPT"

        headers = [
            "STT", "Mã Trường Ghép", "Mã Tỉnh", "Tên Tỉnh/Thành", 
            "Mã Phường Xã", "Tên Xã/Phường", "Mã Trường", 
            "Tên Trường THPT", "Địa Chỉ", "Khu Vực"
        ]
        ws.append(headers)

        danh_sach = TruongTHPT.objects.all().order_by('id')
        for idx, item in enumerate(danh_sach, start=1):
            ws.append([
                idx,
                item.ma_truong_ghep or '',
                item.ma_tinh or '',
                item.ten_tinh or '',
                item.ma_phuong_xa or '',
                item.ten_xa_phuong or '',
                item.ma_truong or '',
                item.ten_truong or '',
                item.dia_chi or '',
                item.khu_vuc or ''
            ])

        wb.save(response)
        return response

    except Exception as e:
        messages.error(request, f"Lỗi khi xuất file Excel: {e}")
        return redirect('danh_sach_truong_thpt')


@custom_login_required
@check_permission('sua_truong_thpt')
def sua_truong_thpt(request, id):
    if request.method == 'POST':
        try:
            truong = TruongTHPT.objects.get(id=id)
            truong.ma_tinh = request.POST.get('ma_tinh', '').strip()
            truong.ma_truong = request.POST.get('ma_truong', '').strip()
            truong.ten_truong = request.POST.get('ten_truong', '').strip()
            truong.khu_vuc = request.POST.get('khu_vuc', '').strip()
            truong.dia_chi = request.POST.get('dia_chi', '').strip()
            truong.ten_tinh = request.POST.get('ten_tinh', '').strip()
            truong.ten_xa_phuong = request.POST.get('ten_xa_phuong', '').strip()

            if truong.ma_tinh and truong.ma_truong:
                truong.ma_truong_ghep = f"{truong.ma_tinh.zfill(2)}{truong.ma_truong.zfill(3)}"

            truong.save()
            messages.success(request, f"Cập nhật thông tin trường {truong.ten_truong} thành công!")
        except TruongTHPT.DoesNotExist:
            messages.error(request, "Không tìm thấy thông tin trường THPT cần sửa!")
        except Exception as e:
            messages.error(request, f"Lỗi khi cập nhật: {str(e)}")

    return redirect('danh_sach_truong_thpt')


@custom_login_required
@check_permission('xoa_truong_thpt')
def xoa_truong_thpt(request, id):
    truong = get_object_or_404(TruongTHPT, id=id)
    if request.method == 'POST':
        ten_truong = truong.ten_truong
        truong.delete()
        messages.success(request, f"Đã xóa thành công trường: {ten_truong}")
    return redirect('danh_sach_truong_thpt')
    
    
def map_diem_vsat_theo_to_hop(vsat_obj, ma_to_hop):
    """Hàm bổ trợ: Ánh xạ điểm 3 môn dựa theo mã tổ hợp."""
    if not vsat_obj or not ma_to_hop:
        return None, None, None
    
    to_hop = str(ma_to_hop).strip().upper()
    mapping = {
        'A00': ('to_vs', 'li_vs', 'ho_vs'),
        'A01': ('to_vs', 'li_vs', 'n1_vs'),
        'B00': ('to_vs', 'ho_vs', 'si_vs'),
        'C00': ('va_vs', 'su_vs', 'di_vs'),
        'D01': ('to_vs', 'va_vs', 'n1_vs'),
        'D07': ('to_vs', 'ho_vs', 'n1_vs'),
        'D09': ('to_vs', 'su_vs', 'n1_vs'),
        'D14': ('va_vs', 'su_vs', 'n1_vs'),
    }
    
    if to_hop in mapping:
        f1, f2, f3 = mapping[to_hop]
        return getattr(vsat_obj, f1, None), getattr(vsat_obj, f2, None), getattr(vsat_obj, f3, None)
    
    return None, None, None


@custom_login_required
@check_permission('import_diem_vsat')
def import_diem_vsat(request):
    """Hàm xử lý riêng cho Import Excel điểm V-SAT"""
    if request.method == 'POST' and request.FILES.get('excel_file'):
        file_excel = request.FILES['excel_file']
        try:
            df = pd.read_excel(file_excel, sheet_name=0)

            def parse_float(val):
                try:
                    if pd.isna(val) or str(val).strip().lower() in ['', 'nan', 'none']:
                        return None
                    return float(val)
                except:
                    return None

            danh_sach_vsat = []
            for _, row in df.iterrows():
                raw_cccd = str(row.iloc[0]).strip().split('.')[0]
                if not raw_cccd or raw_cccd.lower() in ['nan', 'none', '']:
                    continue

                cccd = raw_cccd.zfill(12) if len(raw_cccd) < 12 and raw_cccd.isdigit() else raw_cccd
                ho_ten = str(row.iloc[1]).strip()

                danh_sach_vsat.append(DiemThiVsat(
                    so_cccd=cccd,
                    ho_ten=ho_ten,
                    di_vs=parse_float(row.iloc[2]) if len(row) > 2 else None,
                    ho_vs=parse_float(row.iloc[3]) if len(row) > 3 else None,
                    li_vs=parse_float(row.iloc[4]) if len(row) > 4 else None,
                    n1_vs=parse_float(row.iloc[5]) if len(row) > 5 else None,
                    si_vs=parse_float(row.iloc[6]) if len(row) > 6 else None,
                    su_vs=parse_float(row.iloc[7]) if len(row) > 7 else None,
                    to_vs=parse_float(row.iloc[8]) if len(row) > 8 else None,
                    va_vs=parse_float(row.iloc[9]) if len(row) > 9 else None,
                ))

            if danh_sach_vsat:
                DiemThiVsat.objects.bulk_create(
                    danh_sach_vsat,
                    update_conflicts=True,
                    update_fields=['ho_ten', 'di_vs', 'ho_vs', 'li_vs', 'n1_vs', 'si_vs', 'su_vs', 'to_vs', 'va_vs'],
                    unique_fields=['so_cccd']
                )
                messages.success(request, f"Đã import thành công {len(danh_sach_vsat)} bản ghi điểm V-SAT!")
        except Exception as e:
            messages.error(request, f"Lỗi khi import điểm V-SAT: {str(e)}")

        return redirect(request.META.get('HTTP_REFERER', '/'))

    return redirect('index')



# ==========================================
# VIEW TRA CỨU & XEM GIẤY BÁO DÀNH CHO SINH VIÊN
# ==========================================
def in_giay_bao_bao_mat(request):
    """Khai báo view trùng tên với URL 'in_giay_bao_bao_mat' để hiển thị Loading & PDF"""
    cau_hinh = CauHinhGiayBao.get_config()
    now = timezone.now()

    # 1. Kiểm tra thời gian bắt đầu
    if cau_hinh.ngay_bat_dau and now < cau_hinh.ngay_bat_dau:
        start_str = cau_hinh.ngay_bat_dau.strftime('%H:%M %d/%m/%Y')
        return HttpResponse(
            f'Cổng tra cứu giấy báo chưa mở. Thời gian bắt đầu: {start_str}',
            status=403,
        )

    # 2. Kiểm tra thời gian kết thúc
    if cau_hinh.ngay_ket_thuc and now > cau_hinh.ngay_ket_thuc:
        return HttpResponse(
            'Hệ thống đã đóng cổng tra cứu và in giấy báo trúng tuyển!',
            status=403,
        )

    # 3. Kiểm tra phiên làm việc sinh viên
    student_cccd = request.session.get('student_cccd')
    if not student_cccd:
        return HttpResponse(
            'Phiên làm việc hết hạn hoặc bạn chưa đăng nhập!', status=403
        )

    # 4. Chuyển sang hàm dùng chung (Hàm này tự động trả về màn hình Loading dạng Card)
    return xu_ly_xuat_giay_bao(request, student_cccd)

def sinh_vien_login(request):
    """Trang đăng nhập tra cứu cho sinh viên"""
    if request.method == 'POST':
        identifier = request.POST.get('identifier', '').strip()
        password = request.POST.get('password', '').strip()

        ts = MauImportGiayBao.objects.filter(
            Q(cccd=identifier) | Q(ma_dkxt=identifier) | Q(ma_sv=identifier)
        ).first()

        if ts:
            is_valid_pass = (getattr(ts, 'mat_khau', None) and ts.mat_khau == password)
            is_valid_cccd = (ts.cccd == password)

            if is_valid_pass or is_valid_cccd:
                request.session['student_cccd'] = ts.cccd
                return redirect('sinh_vien_dashboard')

        messages.error(request, 'Thông tin Mã hồ sơ/SBD hoặc Mật khẩu/CCCD không chính xác!')

    return render(request, 'sinhvien/sv_login.html')

def sinh_vien_logout(request):
    """Xóa session và đăng xuất sinh viên"""
    request.session.flush()
    return redirect('sinh_vien_login')

def sinh_vien_dashboard(request):
    student_cccd = request.session.get('student_cccd')
    if not student_cccd:
        return redirect('sinh_vien_login')

    ts = MauImportGiayBao.objects.filter(cccd=student_cccd).first()
    if not ts:
        ts = get_object_or_404(MauImportGiayBao, cccd=student_cccd)

    ts_data = ThiSinhData.objects.filter(cccd=student_cccd).first()

    # Ưu tiên lấy giới tính từ CapNhatThongTinTrungTuyen, fallback sang MauImportGiayBao
    cap_nhat_info = CapNhatThongTinTrungTuyen.objects.filter(cccd=student_cccd).first()
    gioi_tinh = "—"
    if cap_nhat_info and cap_nhat_info.gioi_tinh:
        gioi_tinh = cap_nhat_info.gioi_tinh
    elif hasattr(ts, 'gioi_tinh') and ts.gioi_tinh:
        gioi_tinh = ts.gioi_tinh

    nv_trung_tuyen = "—"
    if ts:
        nv_trung_tuyen = getattr(ts, 'nv_TT', None) or getattr(ts, 'nv_tt', None) or getattr(ts, 'nv_trung_tuyen', None)

    if not nv_trung_tuyen and ts_data:
        nv_trung_tuyen = getattr(ts_data, 'nv_TT', None) or getattr(ts_data, 'nv_tt', None) or getattr(ts_data, 'nv_trung_tuyen', None)

    nv_trung_tuyen = str(nv_trung_tuyen).strip() if (nv_trung_tuyen and str(nv_trung_tuyen).lower() != 'none') else "1"

    raw_ctdt = str(ts.ctdt or '').strip()
    ma_nganh = ""
    ten_nganh = raw_ctdt
    match = re.match(r'^(\d+)[\s:\-–—]*(.*)$', raw_ctdt)
    if match:
        ma_nganh = match.group(1)
        ten_nganh = match.group(2) if match.group(2) else raw_ctdt

    # Xử lý điểm ưu tiên khu vực / đối tượng
    diem_utkv_giam_tru = "—"
    raw_utkv = getattr(ts_data, 'diem_utkv_giam_tru', None) if ts_data else None
    if raw_utkv is not None and raw_utkv != "":
        try:
            diem_utkv_giam_tru = round(float(raw_utkv), 2)
        except (ValueError, TypeError):
            diem_utkv_giam_tru = raw_utkv

    diem_utkv_chua_giam_tru = "—"
    raw_utkv_chua_tru = getattr(ts_data, 'diem_utkv_chua_giam_tru', None) if ts_data else None
    if raw_utkv_chua_tru is not None and raw_utkv_chua_tru != "":
        try:
            diem_utkv_chua_giam_tru = round(float(raw_utkv_chua_tru), 2)
        except (ValueError, TypeError):
            diem_utkv_chua_giam_tru = raw_utkv_chua_tru

    ptxt_upper = str(getattr(ts, 'ptxt', '') or '').upper()
    is_dgnl = ('DGNL' in ptxt_upper) or ('ĐGNL' in ptxt_upper) or ('EVAL' in ptxt_upper)
    is_vsat = ('VSAT' in ptxt_upper) or ('V-SAT' in ptxt_upper)

    if is_dgnl:
        thm_display = None
    elif is_vsat:
        thm_display = getattr(ts, 'pt2_vsat_thm', None) or getattr(ts, 'pt2_tn_thm', None)
    else:
        thm_display = getattr(ts, 'pt2_tn_thm', None)

    m1, m2, m3 = "—", "—", "—"
    diem_dgnl_goc = "—"
    tong_diem_goc = "—"

    if is_dgnl:
        diem_dgnl_val = None
        if ts_data:
            diem_dgnl_val = (
                getattr(ts_data, 'diem_DGNL', None) or 
                getattr(ts_data, 'diem_dgnl', None) or 
                getattr(ts_data, 'diem_max', None)
            )
        if not diem_dgnl_val:
            diem_dgnl_val = getattr(ts, 'diem_DGNL', None) or getattr(ts, 'diem_dgnl', "—") or "—"
        
        diem_dgnl_goc = diem_dgnl_val
        tong_diem_goc = diem_dgnl_val

    elif is_vsat:
        m1 = getattr(ts, 'pt2_vsat_diemmon1', None) or getattr(ts, 'pt2_tn_diemmon1', None) or "—"
        m2 = getattr(ts, 'pt2_vsat_diemmon2', None) or getattr(ts, 'pt2_tn_diemmon2', None) or "—"
        m3 = getattr(ts, 'pt2_vsat_diemmon3', None) or getattr(ts, 'pt2_tn_diemmon3', None) or "—"
        try:
            if m1 != "—" and m2 != "—" and m3 != "—":
                tong_diem_goc = round(float(m1) + float(m2) + float(m3), 2)
        except (ValueError, TypeError):
            tong_diem_goc = "—"

    else:
        m1 = getattr(ts, 'pt2_tn_diemmon1', None) or "—"
        m2 = getattr(ts, 'pt2_tn_diemmon2', None) or "—"
        m3 = getattr(ts, 'pt2_tn_diemmon3', None) or "—"
        try:
            if m1 != "—" and m2 != "—" and m3 != "—":
                tong_diem_goc = round(float(m1) + float(m2) + float(m3), 2)
        except (ValueError, TypeError):
            tong_diem_goc = "—"

    diem_thi_qd = "—"
    diem_thi_60 = "—"
    diem_hb_qd = "—"
    diem_hb_40 = "—"

    if ts_data:
        if getattr(ts_data, 'diem_max', None) is not None:
            try:
                diem_thi_qd = round(float(ts_data.diem_max) * 100 / 30, 2)
            except (ValueError, TypeError):
                pass

        diem_thi_60 = getattr(ts_data, 'diem_thi_max_qd100_nhan_06', "—")

        if getattr(ts_data, 'diem_tb_cac_nam_hoc', None) is not None:
            try:
                diem_hb_qd = round(float(ts_data.diem_tb_cac_nam_hoc) * 10, 2)
            except (ValueError, TypeError):
                pass

        diem_hb_40 = getattr(ts_data, 'diem_hoc_ba_qd100_nhan_04', "—")

    try:
        dx_thuong = float(getattr(ts, 'diem_xet_thuong', 0) or 0)
        c_ccta = float(getattr(ts, 'cong_ccta', 0) or 0)
        diem_thuong_kk = round(dx_thuong + c_ccta, 2)
    except (ValueError, TypeError):
        diem_thuong_kk = 0

    diem_chuan = getattr(ts, 'dc', "—") or "—"
    

    context = {
        'ts': ts,
        'cau_hinh': CauHinhGiayBao.get_config(),
        'gioi_tinh': gioi_tinh,
        'ma_nganh': ma_nganh or "—",
        'ten_nganh': ten_nganh or "—",
        'is_dgnl': is_dgnl,
        'is_vsat': is_vsat,
        'thm_display': thm_display,
        'm1': m1,
        'm2': m2,
        'm3': m3,
        'diem_dgnl_goc': diem_dgnl_goc,
        'tong_diem_goc': tong_diem_goc,
        'diem_thi_qd': diem_thi_qd,
        'diem_thi_60': diem_thi_60,
        'diem_hb_qd': diem_hb_qd,
        'diem_hb_40': diem_hb_40,
        'diem_thuong_kk': diem_thuong_kk,
        'diem_utkv_chua_giam_tru': diem_utkv_chua_giam_tru,
        'diem_utkv_giam_tru': diem_utkv_giam_tru,
        'diem_chuan': diem_chuan,
        'nv_trung_tuyen': nv_trung_tuyen,
    }

    return render(request, 'sinhvien/sv_dashboard.html', context)


